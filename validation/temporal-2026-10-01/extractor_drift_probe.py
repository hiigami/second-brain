"""P12: extractor output changes while the original .docx bytes and TOOL_VERSION stay the same."""
import json, shutil, sys, tempfile, subprocess
from pathlib import Path
ROOT = Path(sys.argv[1]); sys.path.insert(0, str(ROOT/"tools")); sys.path.insert(0, str(ROOT/"tests"))
import docx
from kb_common import read_json, write_json_new
from kb_inventory import inventory
from kb_reanchor import reanchor
from helpers import SOURCES, PROJECT_CONFIG
t = Path(tempfile.mkdtemp()).resolve(); shutil.copytree(SOURCES, t/"sources")
d = docx.Document(); d.add_paragraph("Acta 2026-09-25"); d.add_paragraph("Se acordó usar un botón de descarga."); d.save(t/"sources/meetings/acta.docx")
proj = t/"projects/demo"; (proj/"config").mkdir(parents=True)
cfg = read_json(PROJECT_CONFIG)
for s in cfg["sources"]:
    s["path"] = str(t/"sources"/s["id"])
    if s["id"] == "meetings": s["include"].append("**/*.docx")
(proj/"config/project.json").write_text(json.dumps(cfg)); c = proj/"config/project.json"
r1, m1 = inventory(c, "doc-1", documents=True)
f1 = [f for f in m1["files"] if f["relative_path"] == "acta.docx"][0]
lines = (r1/f1["snapshot_path"]).read_text().splitlines()
n = next(i for i, l in enumerate(lines, 1) if "botón" in l)
rec = {"schema_version": "0.2", "project_id": m1["project_id"], "run_id": "doc-1",
       "coverage": [{"evidence_id": f["evidence_id"], "disposition": "used" if f is f1 else "reviewed_no_record", "note": "probe"} for f in m1["files"]],
       "records": [{"id": "DEC-001", "kind": "decision", "title": "t", "statement": "s", "epistemic_status": "observed",
                    "evidence": [{"evidence_id": f1["evidence_id"], "start_line": n, "end_line": n, "quote": lines[n-1]}],
                    "relations": [], "open_questions": [], "investigation": None}]}
write_json_new(r1/"proposals/records.json", rec)
# Simulate an extractor bugfix: one extra header line, same VERSION, same TOOL_VERSION.
ex = ROOT/"tools/kb_document_extractors.py"; s = ex.read_text()
ex.write_text(s.replace('"NOTICE: Locator labels and this header are generated metadata, not source assertions.", "",',
                        '"NOTICE: Locator labels and this header are generated metadata, not source assertions.", "NOTICE: extra line.", "",', 1))
r2, m2 = inventory(c, "doc-2", baseline=r1, documents=True)
f2 = [f for f in m2["files"] if f["relative_path"] == "acta.docx"][0]
print("original sha same:", f1["sha256"] == f2["sha256"], "| text sha same:", f1["document"]["text_sha256"] == f2["document"]["text_sha256"])
print("evidence id same:", f1["evidence_id"] == f2["evidence_id"], "| delta compatible:", m2["delta"]["compatible"],
      "| acta in unchanged:", f1["logical_id"] in m2["delta"]["unchanged"])
_, rep = reanchor(r1, r2)
print("reanchor:", rep["citation_status_counts"], "| structural_check:", rep["structural_check"]["status"], rep["structural_check"].get("error", ""))
