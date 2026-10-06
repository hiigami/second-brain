"""Synthetic DOCX annotations; no Word export interoperability claim."""
from dataclasses import replace
from io import BytesIO
import json
import sys
from pathlib import Path
import unittest
from xml.etree import ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from kb_document_extractors import ExtractionError, Limits, extract_bytes

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
CID = "http://schemas.microsoft.com/office/word/2016/wordml/cid"
CEX = "http://schemas.microsoft.com/office/word/2018/wordml/cex"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
ROLES = {
    "extended": ("commentsExtended.xml", W15, "commentsEx", "commentEx",
                 "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"),
    "ids": ("commentsIds.xml", CID, "commentsIds", "commentId",
            "http://schemas.microsoft.com/office/2016/09/relationships/commentsIds"),
    "extensible": ("commentsExtensible.xml", CEX, "commentsExtensible", "commentExtensible",
                   "http://schemas.microsoft.com/office/2018/08/relationships/commentsExtensible"),
}


def comment_docx(*, metadata=None, body=True, change=None):
    """Two synthetic comments whose *last* paragraphs identify their thread."""
    from docx import Document
    document = Document()
    if body:
        run = document.add_paragraph("Pending scope.").runs[0]
        first = document.add_comment(run, "First line\nLast line", author="Reviewer", initials="RV")
    else:
        first = document.comments.add_comment("First line\nLast line", author="Reviewer", initials="RV")
    first.add_table(rows=1, cols=1, width=1000000).cell(0, 0).text = "Comment table"
    first.add_paragraph("Final paragraph")
    document.comments.add_comment("Reply text", author="Responder", initials=None)
    stream = BytesIO()
    document.save(stream)
    with zipfile.ZipFile(BytesIO(stream.getvalue())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    comments = ET.fromstring(parts["word/comments.xml"])
    for index, comment in enumerate(comments):
        comment.set(f"{{{W}}}date", "2026-09-01T12:30:00" if index == 0 else "2026-09-02T10:00:00+02:00")
        paragraphs = list(comment.iter(f"{{{W}}}p"))
        paragraphs[0].set(f"{{{W14}}}paraId", f"000000{index + 5:02X}")
        paragraphs[-1].set(f"{{{W14}}}paraId", f"000000{index + 1:02X}")
    parts["word/comments.xml"] = ET.tostring(comments)
    rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
    content_types = ET.fromstring(parts["[Content_Types].xml"])
    for role, attributes in (metadata or {}).items():
        name, namespace, root_name, child_name, relationship_type = ROLES[role]
        root = ET.Element(f"{{{namespace}}}{root_name}")
        for attrs in attributes:
            ET.SubElement(root, f"{{{namespace}}}{child_name}", {f"{{{namespace}}}{k}": v for k, v in attrs.items()})
        parts["word/" + name] = ET.tostring(root)
        ET.SubElement(rels, f"{{{REL}}}Relationship", Id="comment-" + role, Target=name, Type=relationship_type)
        ET.SubElement(content_types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                      PartName="/word/" + name,
                      ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml." + name[:-4] + "+xml")
    parts["word/_rels/document.xml.rels"] = ET.tostring(rels)
    parts["[Content_Types].xml"] = ET.tostring(content_types)
    if change:
        change(parts)
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return stream.getvalue()


def annotation_metadata(result):
    lines = result.text.splitlines()
    return [json.loads("\n".join(lines[s["line_start"] - 1:s["line_end"]]))
            for s in result.metadata["segments"] if s["locator"].endswith("/metadata")]


class DocxCommentTests(unittest.TestCase):
    def test_classic_text_tables_raw_dates_and_unknown_state(self):
        result = extract_bytes(comment_docx(), "comments.docx")
        comments = annotation_metadata(result)
        self.assertEqual(len(comments), 2)
        self.assertEqual(comments[0]["author"], "Reviewer")
        self.assertEqual(comments[0]["initials"], "RV")
        self.assertEqual(comments[0]["date"], "2026-09-01T12:30:00")
        self.assertEqual(comments[1]["date"], "2026-09-02T10:00:00+02:00")
        for comment in comments:
            self.assertIsNone(comment["resolved"])
            self.assertIsNone(comment["parent_comment_id"])
            self.assertIsNone(comment["resolved_at"])
        self.assertIn("Comment table", result.text)
        self.assertIn("First line", result.text)
        self.assertIn("Last line", result.text)
        self.assertIn("Reply text", result.text)
        self.assertTrue(comments[0]["anchors"])

    def test_reply_joins_last_paragraph_and_does_not_inherit_status(self):
        result = extract_bytes(comment_docx(metadata={"extended": [
            {"paraId": "00000001", "done": "1"},
            {"paraId": "00000002", "paraIdParent": "00000001"}]}), "thread.docx")
        first, reply = annotation_metadata(result)
        self.assertTrue(first["resolved"])
        self.assertFalse(reply["resolved"])
        self.assertEqual(reply["parent_comment_id"], first["comment_id"])
        self.assertEqual(reply["status_basis"], "commentEx/default_done_false")
        self.assertIsNone(first["resolved_at"])

    def test_all_on_off_forms(self):
        for flag, value in (("true", True), ("1", True), ("on", True), ("false", False), ("0", False), ("off", False)):
            with self.subTest(flag=flag):
                result = extract_bytes(comment_docx(metadata={"extended": [{"paraId": "00000001", "done": flag}]}), "flag.docx")
                self.assertIs(annotation_metadata(result)[0]["resolved"], value)

    def test_missing_and_empty_fields_remain_distinct(self):
        def change(parts):
            root = ET.fromstring(parts["word/comments.xml"])
            root[0].attrib.pop(f"{{{W}}}date")
            root[0].set(f"{{{W}}}author", "")
            root[0].set(f"{{{W}}}initials", "")
            parts["word/comments.xml"] = ET.tostring(root)
        first, reply = annotation_metadata(extract_bytes(comment_docx(change=change), "empty.docx"))
        self.assertEqual(first["author"], "")
        self.assertEqual(first["initials"], "")
        self.assertIsNone(first["date"])
        self.assertIsNone(reply["initials"])

    def test_modern_date_and_placeholder_are_not_resolution_history(self):
        result = extract_bytes(comment_docx(metadata={
            "ids": [{"paraId": "00000001", "durableId": "0000000A"}],
            "extensible": [{"durableId": "0000000A", "dateUtc": "2026-09-01T12:30:00Z", "intelligentPlaceholder": "true"}]}), "modern.docx")
        first = annotation_metadata(result)[0]
        self.assertEqual(first["date_utc"], "2026-09-01T12:30:00Z")
        self.assertTrue(first["intelligent_placeholder"])
        self.assertIsNone(first["resolved_at"])
        self.assertNotIn("First line", result.text)
        self.assertNotIn("Comment table", result.text)
        self.assertIn("Reply text", result.text)
        self.assertIn("docx_comment_placeholder_body_unavailable", [i["code"] for i in result.metadata["issues"]])

    def test_malformed_thread_metadata_is_rejected(self):
        cases = [
            [{"paraId": "00000001", "done": "yes"}],
            [{"paraId": "00000001"}, {"paraId": "00000001"}],
            [{"paraId": "00000003"}],
            [{"paraId": "00000001", "paraIdParent": "00000003"}],
            [{"paraId": "00000001", "paraIdParent": "00000001"}],
            [{"paraId": "00000001", "paraIdParent": "00000002"}, {"paraId": "00000002", "paraIdParent": "00000001"}],
        ]
        for attributes in cases:
            with self.subTest(attributes=attributes):
                with self.assertRaises(ExtractionError) as caught:
                    extract_bytes(comment_docx(metadata={"extended": attributes}), "bad.docx")
                self.assertEqual(caught.exception.code, "docx_comment_metadata_invalid")

    def test_comment_only_document_is_extractable(self):
        result = extract_bytes(comment_docx(body=False), "only-comments.docx")
        self.assertIn("Reply text", result.text)
        self.assertEqual(len(annotation_metadata(result)), 2)

    def test_duplicate_comment_and_paragraph_ids_are_rejected(self):
        for mode in ("comment", "paragraph"):
            def change(parts):
                root = ET.fromstring(parts["word/comments.xml"])
                if mode == "comment":
                    root[1].set(f"{{{W}}}id", "00")
                else:
                    list(root[1].iter(f"{{{W}}}p"))[-1].set(f"{{{W14}}}paraId", "00000001")
                parts["word/comments.xml"] = ET.tostring(root)
            with self.subTest(mode=mode):
                with self.assertRaises(ExtractionError) as caught:
                    extract_bytes(comment_docx(change=change), "duplicate.docx")
                self.assertEqual(caught.exception.code, "docx_comment_metadata_invalid")

    def test_durable_id_joins_are_one_to_one(self):
        cases = [
            {"ids": [{"paraId": "00000003", "durableId": "0000000A"}]},
            {"ids": [{"paraId": "00000001", "durableId": "0000000A"}, {"paraId": "00000002", "durableId": "0000000A"}]},
            {"extensible": [{"durableId": "0000000A"}]},
            {"ids": [{"paraId": "00000001", "durableId": "00000000"}]},
            {"ids": [{"paraId": "00000001", "durableId": "0000000A"}],
             "extensible": [{"durableId": "0000000A"}, {"durableId": "0000000A"}]},
        ]
        for metadata in cases:
            with self.subTest(metadata=metadata):
                with self.assertRaises(ExtractionError) as caught:
                    extract_bytes(comment_docx(metadata=metadata), "ids.docx")
                self.assertEqual(caught.exception.code, "docx_comment_metadata_invalid")

    def test_placeholder_attribute_on_reply_is_rejected(self):
        for flag in ("true", "false"):
            metadata = {"extended": [{"paraId": "00000002", "paraIdParent": "00000001"}],
                        "ids": [{"paraId": "00000002", "durableId": "0000000A"}],
                        "extensible": [{"durableId": "0000000A", "intelligentPlaceholder": flag}]}
            with self.subTest(flag=flag):
                with self.assertRaises(ExtractionError) as caught:
                    extract_bytes(comment_docx(metadata=metadata), "reply.docx")
                self.assertEqual(caught.exception.code, "docx_comment_metadata_invalid")

    def test_nonstandard_relationship_selected_comment_part(self):
        def change(parts):
            parts["annotations/review.xml"] = parts.pop("word/comments.xml")
            rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
            for rel in rels:
                if rel.get("Type").endswith("/comments"):
                    rel.set("Target", "../annotations/review.xml")
            parts["word/_rels/document.xml.rels"] = ET.tostring(rels)
            types = ET.fromstring(parts["[Content_Types].xml"])
            for item in types:
                if item.get("PartName") == "/word/comments.xml":
                    item.set("PartName", "/annotations/review.xml")
            parts["[Content_Types].xml"] = ET.tostring(types)
        result = extract_bytes(comment_docx(change=change), "custom.docx")
        self.assertIn("Reply text", result.text)
        self.assertTrue(any(s["locator"] == "DOCX/part=annotations/review.xml/comment=0/metadata" for s in result.metadata["segments"]))

    def test_ambiguous_missing_external_and_invalid_part_containers(self):
        for mode in ("duplicate", "missing", "external", "wrong-root", "wrong-namespace", "duplicate-relationship-id"):
            def change(parts):
                rels = ET.fromstring(parts["word/_rels/document.xml.rels"])
                rel = next(r for r in rels if r.get("Type").endswith("/comments"))
                if mode == "duplicate":
                    ET.SubElement(rels, rel.tag, dict(rel.attrib, Id="duplicate"))
                elif mode == "missing":
                    rel.set("Target", "absent.xml")
                elif mode == "external":
                    rel.set("TargetMode", "External")
                elif mode == "duplicate-relationship-id":
                    rel.set("Id", rels[0].get("Id"))
                else:
                    root = ET.fromstring(parts["word/comments.xml"])
                    root.tag = f"{{{W}}}wrong" if mode == "wrong-root" else "{urn:unknown}comments"
                    parts["word/comments.xml"] = ET.tostring(root)
                parts["word/_rels/document.xml.rels"] = ET.tostring(rels)
            with self.subTest(mode=mode):
                with self.assertRaises(ExtractionError) as caught:
                    extract_bytes(comment_docx(change=change), "part.docx")
                self.assertEqual(caught.exception.code, "docx_comment_metadata_invalid")

    def test_unknown_extensions_and_unlinked_parts_are_visible(self):
        def change(parts):
            root = ET.fromstring(parts["word/commentsExtended.xml"])
            ET.SubElement(root[0], f"{{{W15}}}extLst")
            parts["word/commentsExtended.xml"] = ET.tostring(root)
            parts["word/orphan.xml"] = f'<comments xmlns="{W}"/>'.encode()
        result = extract_bytes(comment_docx(metadata={"extended": [{"paraId": "00000001"}]}, change=change), "extensions.docx")
        codes = {issue["code"] for issue in result.metadata["issues"]}
        self.assertTrue({"docx_comment_extensions_unavailable", "docx_comment_parts_unavailable"} <= codes)
        self.assertIn("Reply text", result.text)

    def test_structure_and_output_budgets_reject_without_truncation(self):
        limits = Limits()
        for bounded in (replace(limits, max_structure_nodes=5), replace(limits, max_structure_depth=4),
                        replace(limits, max_output_chars=500), replace(limits, max_segments=2)):
            with self.subTest(limits=bounded):
                with self.assertRaises(ExtractionError) as caught:
                    extract_bytes(comment_docx(), "budget.docx", limits=bounded)
                self.assertEqual(caught.exception.code, "extraction_budget_exceeded")

    def test_hyperlink_text_is_retained_and_external_target_is_inert(self):
        def change(parts):
            root = ET.fromstring(parts["word/comments.xml"])
            paragraph = list(root[0].iter(f"{{{W}}}p"))[0]
            link = ET.SubElement(paragraph, f"{{{W}}}hyperlink",
                                 {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": "external"})
            ET.SubElement(ET.SubElement(link, f"{{{W}}}r"), f"{{{W}}}t").text = "Linked annotation"
            parts["word/comments.xml"] = ET.tostring(root)
            parts["word/_rels/comments.xml.rels"] = (f'<Relationships xmlns="{REL}"><Relationship Id="external" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
                'Target="https://invalid.example.test/do-not-fetch" TargetMode="External"/></Relationships>').encode()
        from unittest.mock import patch
        with patch("urllib.request.urlopen", side_effect=AssertionError("Network access is forbidden")):
            result = extract_bytes(comment_docx(change=change), "link.docx")
        self.assertIn("Linked annotation", result.text)
        self.assertIn("external_relationships_not_fetched", {i["code"] for i in result.metadata["issues"]})

    def test_malformed_dates_and_modern_flags_are_rejected(self):
        for date in ("invalid", "2026-02-30T12:30:00", "2026-09-01T12:30:00+01:99"):
            def change(parts):
                root = ET.fromstring(parts["word/comments.xml"])
                root[0].set(f"{{{W}}}date", date)
                parts["word/comments.xml"] = ET.tostring(root)
            with self.subTest(date=date), self.assertRaises(ExtractionError):
                extract_bytes(comment_docx(change=change), "date.docx")
        for attrs in ({"dateUtc": "2026-09-01T12:30:00+02:00"}, {"intelligentPlaceholder": "yes"}):
            with self.subTest(attrs=attrs), self.assertRaises(ExtractionError):
                extract_bytes(comment_docx(metadata={"ids": [{"paraId": "00000001", "durableId": "0000000A"}],
                    "extensible": [dict(attrs, durableId="0000000A")]}), "modern.docx")

    def test_date_utc_without_lexical_timezone_retains_attribute_semantics(self):
        date = "2026-09-01T12:30:00"
        result = extract_bytes(comment_docx(metadata={"ids": [{"paraId": "00000001", "durableId": "0000000A"}],
            "extensible": [{"durableId": "0000000A", "dateUtc": date}]}), "implicit-utc.docx")
        self.assertEqual(annotation_metadata(result)[0]["date_utc"], date)
        self.assertIsNone(annotation_metadata(result)[0]["resolved_at"])

    def test_unsupported_inline_text_is_disclosed_as_unavailable(self):
        for wrapper in ("customXml", "smartTag", "fldSimple"):
            def change(parts):
                root = ET.fromstring(parts["word/comments.xml"])
                paragraph = next(root[0].iter(f"{{{W}}}p"))
                enclosing = ET.SubElement(paragraph, f"{{{W}}}{wrapper}")
                ET.SubElement(ET.SubElement(enclosing, f"{{{W}}}r"), f"{{{W}}}t").text = "Omitted wrapped prose"
                parts["word/comments.xml"] = ET.tostring(root)
            with self.subTest(wrapper=wrapper):
                result = extract_bytes(comment_docx(change=change), "wrapped.docx")
                self.assertNotIn("Omitted wrapped prose", result.text)
                self.assertIn("docx_comment_body_content_unavailable", {i["code"] for i in result.metadata["issues"]})
                self.assertIn("Reply text", result.text)

    def test_ambiguous_anchors_never_fabricate_target_text(self):
        def change(parts):
            root = ET.fromstring(parts["word/document.xml"])
            paragraph = next(root.iter(f"{{{W}}}p"))
            ET.SubElement(paragraph, f"{{{W}}}commentRangeStart", {f"{{{W}}}id": "0"})
            parts["word/document.xml"] = ET.tostring(root)
        result = extract_bytes(comment_docx(change=change), "anchor.docx")
        self.assertIn("docx_comment_anchor_unavailable", {i["code"] for i in result.metadata["issues"]})
        self.assertNotIn("target_text", annotation_metadata(result)[0])

    def test_document_without_comments_retains_body(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Ordinary paragraph")
        stream = BytesIO()
        document.save(stream)
        result = extract_bytes(stream.getvalue(), "plain.docx")
        self.assertEqual(annotation_metadata(result), [])
        self.assertIn("Ordinary paragraph", result.text)
        self.assertFalse(any(i["code"].startswith("docx_comment_") for i in result.metadata["issues"]))

    def test_comment_markers_without_linked_part_are_disclosed(self):
        marker_names = ("commentReference", "commentRangeStart", "commentRangeEnd")
        for retained in [(name,) for name in marker_names] + [marker_names]:
            for keep_unlinked_part in (False, True):
                def change(parts):
                    rels_name = "word/_rels/document.xml.rels"
                    relationships = ET.fromstring(parts[rels_name])
                    for relationship in list(relationships):
                        if relationship.get("Type", "").endswith("/comments"):
                            relationships.remove(relationship)
                    parts[rels_name] = ET.tostring(relationships)
                    if not keep_unlinked_part:
                        parts.pop("word/comments.xml")
                        types = ET.fromstring(parts["[Content_Types].xml"])
                        for item in list(types):
                            if item.get("PartName") == "/word/comments.xml":
                                types.remove(item)
                        parts["[Content_Types].xml"] = ET.tostring(types)
                    root = ET.fromstring(parts["word/document.xml"])
                    for parent in root.iter():
                        for child in list(parent):
                            if child.tag in {f"{{{W}}}{name}" for name in marker_names if name not in retained}:
                                parent.remove(child)
                    parts["word/document.xml"] = ET.tostring(root)

                with self.subTest(retained=retained, keep_unlinked_part=keep_unlinked_part):
                    result = extract_bytes(comment_docx(change=change), "unavailable-comments.docx")
                    self.assertIn("Pending scope.", result.text)
                    self.assertEqual(annotation_metadata(result), [])
                    issues = [i for i in result.metadata["issues"] if i["code"] == "docx_comment_anchor_unavailable"]
                    self.assertEqual(len(issues), 1)
                    self.assertEqual(issues[0]["locator"], "DOCX/part=word/document.xml")
                    self.assertNotIn("Reply text", result.text)


if __name__ == "__main__":
    unittest.main()
