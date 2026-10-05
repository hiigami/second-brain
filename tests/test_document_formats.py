"""Synthetic adapter and snapshot regressions; not the unavailable Stage 2 suite."""
from dataclasses import replace
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import struct
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import zlib

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from kb_document_extractors import ExtractionError, Limits, Options, digest, extract_bytes
from kb_extract_document import read_stable, verify


def package(parts):
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return stream.getvalue()


def edit_package(raw, name, change):
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        parts = {n: archive.read(n) for n in archive.namelist()}
    parts[name] = change(parts.get(name, b""))
    return package(parts)


def docx_fixture():
    from docx import Document
    document = Document()
    document.add_paragraph("Primero: proveedor aprobado.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Código"
    table.cell(0, 1).text = "00123"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested fact"
    document.add_paragraph("Después de la tabla.")
    document.sections[0].header.paragraphs[0].text = "Header fact"
    document.sections[0].footer.paragraphs[0].text = "Footer fact"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def pptx_fixture():
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = "Proveedor autorizado"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(7), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Código", "00123"
    slide.notes_slide.notes_text_frame.text = "Propuesta, no decisión."
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second._element.set("show", "0")
    group = second.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Grouped text"
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def pdf_fixture(blank=False, mixed=False):
    from reportlab.pdfgen import canvas
    stream = BytesIO()
    pdf = canvas.Canvas(stream)
    if not blank:
        pdf.drawString(72, 720, "Approved supplier 00123")
    pdf.showPage()
    if mixed:
        pdf.showPage()
    pdf.save()
    return stream.getvalue()


def xlsx_fixture():
    n = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    return package({
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        "xl/workbook.xml": f'<workbook {n} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><workbookPr date1904="1"/><sheets><sheet name="Suppliers" sheetId="1" r:id="rId1" state="hidden"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/></Relationships>',
        "xl/sharedStrings.xml": f'<sst {n}><si><t>Código</t></si></sst>',
        "xl/styles.xml": f'<styleSheet {n}><numFmts count="1"><numFmt numFmtId="164" formatCode="00000"/></numFmts><cellXfs count="2"><xf numFmtId="0"/><xf numFmtId="164"/></cellXfs></styleSheet>',
        "xl/worksheets/sheet1.xml": f'''<worksheet {n}><cols><col min="2" max="2" hidden="1"/></cols><sheetData>
<row r="1"><c r="A1" t="s"><v>0</v></c></row>
<row r="2" hidden="1"><c r="A2" t="inlineStr"><is><t>00123</t></is></c>
<c r="B2"><v>12.50</v></c><c r="C2"><f>SUM(B2,1)</f><v>13.50</v></c>
<c r="D2"><f>SUM(B2,2)</f><v/></c><c r="E2" s="1"><v>123</v></c>
<c r="F2" t="b"><v>1</v></c></row></sheetData><mergeCells><mergeCell ref="A3:B3"/></mergeCells></worksheet>''',
    })


class AdapterTests(unittest.TestCase):
    def assert_code(self, code, fn, *args, **kwargs):
        with self.assertRaises(ExtractionError) as context:
            fn(*args, **kwargs)
        self.assertEqual(context.exception.code, code)

    def test_svg_labels_and_vector_geometry_remain_separate_evidence(self):
        raw = (ROOT / "tests/fixtures/patch_g_diagram.svg").read_bytes()
        e = extract_bytes(raw, "diagram.svg")
        self.assertIn("Supplier", e.text)
        self.assertIn("Invoice", e.text)
        self.assertIn("1:N", e.text)
        self.assertIn("0.05", e.text)
        self.assertIn("marker-end", e.text)
        self.assertIn("one-to-one", e.text)
        self.assertNotIn("Supplier -> Invoice", e.text)
        self.assertIn("vector_geometry", {s.get("content_type") for s in e.metadata["segments"]})
        self.assertIn("svg_external_resource_not_fetched", {i["code"] for i in e.metadata["issues"]})
        self.assertIn("svg_rendering_unverified", {i["code"] for i in e.metadata["issues"]})

    def test_png_region_and_annotation_do_not_claim_pixel_meaning(self):
        raw = (ROOT / "tests/fixtures/patch_g_pixels.png").read_bytes()
        e = extract_bytes(raw, "pixels.png")
        self.assertIn('"width": 2', e.text)
        self.assertIn('"height": 1', e.text)
        self.assertIn("green state (unverified annotation)", e.text)
        self.assertIn("image_region", {s.get("content_type") for s in e.metadata["segments"]})
        self.assertIn("png_pixels_not_interpreted", {i["code"] for i in e.metadata["issues"]})
        self.assertNotIn("red state", e.text)

    def test_svg_doctype_and_truncated_png_block(self):
        self.assert_code("unsafe_or_invalid_xml", extract_bytes,
                         b'<!DOCTYPE svg [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
                         b'<svg xmlns="http://www.w3.org/2000/svg"><text>&x;</text></svg>', "bad.svg")
        raw = (ROOT / "tests/fixtures/patch_g_pixels.png").read_bytes()
        self.assert_code("invalid_png", extract_bytes, raw[:-4], "bad.png")

    def test_png_crc_and_pixel_budget_fail_closed(self):
        raw = (ROOT / "tests/fixtures/patch_g_pixels.png").read_bytes()
        damaged = bytearray(raw)
        damaged[20] ^= 1
        self.assert_code("invalid_png", extract_bytes, bytes(damaged), "bad.png")
        self.assert_code("extraction_budget_exceeded", extract_bytes, raw, "pixels.png",
                         limits=replace(Limits(), max_image_pixels=1))
        self.assert_code("extraction_budget_exceeded", extract_bytes, raw, "pixels.png",
                         limits=replace(Limits(), max_png_chunks=1))

    def test_png_corrupt_or_oversized_pixel_stream_is_rejected(self):
        raw = (ROOT / "tests/fixtures/patch_g_pixels.png").read_bytes()
        idat = raw.index(b"IDAT")
        old_length = struct.unpack_from(">I", raw, idat - 4)[0]
        def replace_idat(payload):
            chunk = (struct.pack(">I", len(payload)) + b"IDAT" + payload
                     + struct.pack(">I", zlib.crc32(b"IDAT" + payload)))
            return raw[:idat - 4] + chunk + raw[idat + 4 + old_length + 4:]
        self.assert_code("invalid_png", extract_bytes, replace_idat(b"not zlib"), "bad.png")
        self.assert_code("extraction_budget_exceeded", extract_bytes,
                         replace_idat(zlib.compress(b"\x00" * 10_000)), "oversized.png")
        self.assert_code("invalid_png", extract_bytes,
                         replace_idat(zlib.compress(b"\x05\xff\x00\x00\x00\x00\x00")),
                         "bad-filter.png")

    def test_png_interlaced_and_palette_rows_validate_without_visual_claim(self):
        def chunk(kind, data):
            return (struct.pack(">I", len(data)) + kind + data
                    + struct.pack(">I", zlib.crc32(kind + data)))
        signature = b"\x89PNG\r\n\x1a\n"
        rgb = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 1)
        interlaced = signature + chunk(b"IHDR", rgb) + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b"")
        e = extract_bytes(interlaced, "interlaced.png")
        self.assertEqual(e.metadata["details"]["png_scanlines_validated"], 1)
        self.assertFalse(e.metadata["details"]["pixel_values_reconstructed"])
        palette = struct.pack(">IIBBBBB", 1, 1, 8, 3, 0, 0, 0)
        indexed = (signature + chunk(b"IHDR", palette) + chunk(b"PLTE", b"\xff\x00\x00")
                   + chunk(b"IDAT", zlib.compress(b"\x00\x00")) + chunk(b"IEND", b""))
        self.assertEqual(extract_bytes(indexed, "indexed.png").metadata["details"]["png_scanlines_validated"], 1)
        missing_palette = signature + chunk(b"IHDR", palette) + chunk(b"IDAT", zlib.compress(b"\x00\x00")) + chunk(b"IEND", b"")
        self.assert_code("invalid_png", extract_bytes, missing_palette, "missing-palette.png")
        compressed = zlib.compress(b"\x00\xff\x00\x00")
        split = (signature + chunk(b"IHDR", rgb)
                 + chunk(b"IDAT", compressed[:3]) + chunk(b"tEXt", b"note\x00metadata")
                 + chunk(b"IDAT", compressed[3:]) + chunk(b"IEND", b""))
        self.assert_code("invalid_png", extract_bytes, split, "split.png")

    def test_visual_probe_png_keeps_false_annotation_separate_from_pixels(self):
        raw = (ROOT / "tests/fixtures/patch_g_visual_probe.png").read_bytes()
        e = extract_bytes(raw, "visual-probe.png")
        self.assertIn("entire image is green", e.text)
        self.assertEqual({s["content_type"] for s in e.metadata["segments"]},
                         {"image_annotation", "image_region"})
        self.assertEqual(e.metadata["details"]["png_scanlines_validated"], 32)
        self.assertIn("png_annotation_unverified", {i["code"] for i in e.metadata["issues"]})
        self.assertIn("png_pixels_not_interpreted", {i["code"] for i in e.metadata["issues"]})

    def test_svg_budgets_and_active_content_do_not_emit_claims(self):
        raw = b'<svg xmlns="http://www.w3.org/2000/svg"><script><text>Fake</text></script><text>Visible source label</text></svg>'
        e = extract_bytes(raw, "active.svg")
        self.assertIn("Visible source label", e.text)
        self.assertNotIn("Fake", e.text)
        self.assertIn("svg_active_or_foreign_content_not_extracted", {i["code"] for i in e.metadata["issues"]})
        self.assert_code("extraction_budget_exceeded", extract_bytes, raw, "active.svg",
                         limits=replace(Limits(), max_structure_nodes=2))

    def test_inline_svg_exposes_label_and_geometry_with_visual_gap(self):
        raw = b'<html><body><svg viewBox="0 0 20 10"><text x="2" y="5">Pending</text><line x1="1" y1="2" x2="10" y2="2"/></svg></body></html>'
        e = extract_bytes(raw, "inline.html")
        self.assertIn("Pending", e.text)
        self.assertIn("x2", e.text)
        self.assertIn("diagram_label", {s.get("content_type") for s in e.metadata["segments"]})
        self.assertIn("html_visual_content_not_extracted", {i["code"] for i in e.metadata["issues"]})

    def test_dense_visual_reference_preserves_material_text_and_both_visual_gaps(self):
        raw = (ROOT / "tests/fixtures/patch_g_dense_visual.html").read_bytes()
        e = extract_bytes(raw, "dense-visual.html")
        for value in ("Customer", "Order", "1:N", "customer_id", "Integration Alpha",
                      "SELECT customer_id, total", "total added", "2026-09-01", "10.20", "USD",
                      "draft", "approve", "active", "pending", "not approved business changes"):
            with self.subTest(value=value):
                self.assertIn(value, e.text)
        visual_gaps = [i for i in e.metadata["issues"]
                       if i["code"] == "html_visual_content_not_extracted" and "/svg=" in i["locator"]]
        self.assertEqual(len(visual_gaps), 2)
        self.assertEqual(sum(s["content_type"] == "image_region" for s in e.metadata["segments"]), 2)
        self.assertIn("vector_geometry", {s["content_type"] for s in e.metadata["segments"]})

    def test_inline_svg_inside_suppressed_template_has_no_citable_visuals(self):
        raw = b'<html><template><svg viewBox="0 0 20 10"><text>Fake</text><line x2="10"/></svg></template><p>Actual</p></html>'
        e = extract_bytes(raw, "template.html")
        self.assertIn("Actual", e.text)
        self.assertNotIn("Fake", e.text)
        self.assertNotIn('"x2": "10"', e.text)

    def test_suppressed_html_descendants_emit_no_image_or_resource_evidence(self):
        for tag in ("template", "object", "iframe", "canvas", "foreignobject"):
            with self.subTest(tag=tag):
                raw = (f'<html><{tag} data="parent-resource"><div><template>'
                       '<img src="cid:hidden" alt="Hidden assertion">'
                       '<a href="https://invalid.example/hidden">Hidden link</a>'
                       '</template></div>'
                       f'</{tag}><img src="visible.png" alt="Visible assertion">'
                       '<p>Actual</p></html>').encode()
                e = extract_bytes(raw, "suppressed.html")
                self.assertNotIn("Hidden", e.text)
                self.assertNotIn("cid:hidden", e.text)
                self.assertNotIn("invalid.example/hidden", e.text)
                self.assertIn("parent-resource", e.text)
                self.assertIn("Visible assertion", e.text)
                self.assertIn("visible.png", e.text)
                self.assertIn("Actual", e.text)
                self.assertTrue({"html_active_content_not_extracted", "html_visual_content_not_extracted"}
                                & {i["code"] for i in e.metadata["issues"]})

    def test_suppressed_email_html_does_not_reference_hidden_content_ids(self):
        raw = (b"MIME-Version: 1.0\r\nContent-Type: text/html; charset=utf-8\r\n\r\n"
               b'<template><img src="cid:hidden" alt="Hidden assertion"></template>'
               b'<img src="cid:visible" alt="Visible assertion"><p>Actual</p>')
        e = extract_bytes(raw, "suppressed.eml")
        self.assertNotIn("Hidden assertion", e.text)
        self.assertNotIn("cid:hidden", e.text)
        self.assertIn("Visible assertion", e.text)
        missing = [i["locator"] for i in e.metadata["issues"]
                   if i["code"] == "eml_referenced_attachment_missing"]
        self.assertEqual(missing, ["cid:visible"])

    def test_docx_embedded_png_has_asset_locator_and_gap(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Figure caption: pending")
        document.add_picture(BytesIO((ROOT / "tests/fixtures/patch_g_pixels.png").read_bytes()))
        stream = BytesIO()
        document.save(stream)
        e = extract_bytes(stream.getvalue(), "illustrated.docx")
        self.assertIn("word/media/", e.text)
        self.assertIn("embedded_asset", {s.get("content_type") for s in e.metadata["segments"]})
        self.assertIn("docx_embedded_image_not_extracted", {i["code"] for i in e.metadata["issues"]})

    def test_pptx_embedded_png_has_shape_region_and_gap(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(BytesIO((ROOT / "tests/fixtures/patch_g_pixels.png").read_bytes()),
                                 Inches(1), Inches(2), Inches(2), Inches(1))
        stream = BytesIO()
        deck.save(stream)
        e = extract_bytes(stream.getvalue(), "picture.pptx")
        self.assertIn("bbox_emu", e.text)
        self.assertIn("pptx_image_not_extracted", {i["code"] for i in e.metadata["issues"]})

    def test_xlsx_media_part_has_asset_locator_and_gap(self):
        raw = edit_package(xlsx_fixture(), "xl/media/image1.png",
                           lambda _: (ROOT / "tests/fixtures/patch_g_pixels.png").read_bytes())
        e = extract_bytes(raw, "picture.xlsx")
        self.assertIn("xl/media/image1.png", e.text)
        self.assertIn("xlsx_embedded_image_not_extracted", {i["code"] for i in e.metadata["issues"]})

    def test_pdf_embedded_image_has_page_locator_and_gap(self):
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        stream = BytesIO()
        pdf = canvas.Canvas(stream)
        pdf.drawString(72, 720, "Caption: proposed")
        pdf.drawImage(ImageReader(BytesIO((ROOT / "tests/fixtures/patch_g_pixels.png").read_bytes())),
                      72, 650, width=40, height=20)
        pdf.save()
        e = extract_bytes(stream.getvalue(), "image.pdf")
        self.assertIn("PDF/page=1/xobject=", e.text)
        self.assertIn("pdf_embedded_image_not_extracted", {i["code"] for i in e.metadata["issues"]})

    def test_html_minified_structure_is_located_without_active_content(self):
        raw = (ROOT / "tests/fixtures/patch_e_dense.html").read_bytes()
        e = extract_bytes(raw, "model.html")
        self.assertIn("SELECT supplier_id FROM supplier", e.text)
        self.assertIn("https://invalid.example/spec", e.text)
        self.assertIn("HTML/", e.metadata["segments"][0]["locator"])
        self.assertNotIn("window.fetch", e.text)
        self.assertIn("html_active_content_not_extracted", {i["code"] for i in e.metadata["issues"]})
        self.assertIn("html_external_resource_not_fetched", {i["code"] for i in e.metadata["issues"]})

    def test_eml_multipart_preserves_variants_and_attachment_gap(self):
        raw = (ROOT / "tests/fixtures/patch_e_multipart.eml").read_bytes()
        e = extract_bytes(raw, "message.eml")
        self.assertIn("Subject: Supplier model review", e.text)
        self.assertIn("Earlier suggestion", e.text)
        self.assertIn("model.pdf", e.text)
        self.assertNotIn("%PDF", e.text)
        self.assertIn("eml_attachment_not_extracted", {i["code"] for i in e.metadata["issues"]})
        self.assertIn("eml_referenced_attachment_missing", {i["code"] for i in e.metadata["issues"]})

    def test_html_declared_encoding_and_malformed_tags_are_visible(self):
        raw = '<meta charset="windows-1252"><h1>Café</h1><p>pending'.encode("cp1252")
        e = extract_bytes(raw, "legacy.html")
        self.assertIn("Café", e.text)
        self.assertIn("html_malformed_structure", {i["code"] for i in e.metadata["issues"]})
        self.assert_code("charset_decode_error", extract_bytes,
                         b'<meta charset="utf-8"><p>\xe9</p>', "bad.html")
        self.assert_code("unsupported_charset", extract_bytes,
                         b'<meta charset="x-unknown"><p>text</p>', "bad.html")

    def test_html_structure_budget_blocks_without_partial_evidence(self):
        self.assert_code("extraction_budget_exceeded", extract_bytes,
                         b"<p>one</p><p>two</p>", "many.html",
                         limits=replace(Limits(), max_structure_nodes=1))

    def test_html_javascript_url_and_event_handler_are_inert_gaps(self):
        e = extract_bytes(b'<p onclick="run()">Visible</p><a href="javascript:run()">Link</a>',
                          "inert.html")
        self.assertIn("Visible", e.text)
        self.assertIn("html_active_content_not_extracted", {i["code"] for i in e.metadata["issues"]})

    def test_html_inline_code_preserves_surrounding_qualifiers(self):
        for tag in ("p", "li", "blockquote", "pre", "td", "caption", "span"):
            with self.subTest(tag=tag):
                raw = f"<{tag}>Before <code><span>NOT</span></code> approved</{tag}>".encode()

                result = extract_bytes(raw, "qualifier.html")

                lines = result.text.splitlines()
                values = ["\n".join(lines[s["line_start"] - 1:s["line_end"]])
                          for s in result.metadata["segments"]]
                self.assertIn("Before NOT approved", values)
                self.assertNotIn("Before  approved", values)

    def test_eml_html_inline_code_preserves_surrounding_qualifiers(self):
        raw = (b"MIME-Version: 1.0\r\nContent-Type: text/html; charset=utf-8\r\n\r\n"
               b"<p>Before <code>NOT</code> approved</p>")

        result = extract_bytes(raw, "qualifier.eml")

        lines = result.text.splitlines()
        values = ["\n".join(lines[s["line_start"] - 1:s["line_end"]])
                  for s in result.metadata["segments"]]
        self.assertIn("Before NOT approved", values)
        self.assertNotIn("Before  approved", values)

    def test_html_line_breaks_preserve_separate_values(self):
        for markup in (b"<td>10<br>20</td>", b"<td>10<br/>20</td>",
                       b"<td><p>10<br>20</p></td>"):
            with self.subTest(markup=markup):
                raw = b"<p>NOT<br>approved</p><table><tr>" + markup + b"</tr></table>"

                result = extract_bytes(raw, "values.html")

                lines = result.text.splitlines()
                values = ["\n".join(lines[s["line_start"] - 1:s["line_end"]])
                          for s in result.metadata["segments"]]
                self.assertIn("NOT\napproved", values)
                self.assertIn("10\n20", values)
                row = next(s for s in result.metadata["segments"] if s.get("content_type") == "table_row")
                columns = json.loads("\n".join(lines[row["line_start"] - 1:row["line_end"]]))["columns"]
                self.assertEqual(columns[0]["value"], "10\n20")

    def test_html_model_tables_keep_headers_and_row_context(self):
        raw = (ROOT / "tests/fixtures/patch_f_model.html").read_bytes()
        e = extract_bytes(raw, "model.html")
        rows = [s for s in e.metadata["segments"] if s.get("content_type") == "table_row"]
        self.assertEqual(4, len(rows))
        row_lines = e.text.splitlines()[rows[0]["line_start"] - 1:rows[0]["line_end"]]
        row = json.loads("\n".join(row_lines))
        self.assertEqual("Supplier entity and values", row["preceding_heading"])
        self.assertEqual("credit_limit", row["columns"][1]["value"])
        self.assertEqual("Column", row["columns"][1]["header"])
        self.assertEqual("USD", row["columns"][3]["value"])
        self.assertIn("SELECT supplier_id, credit_limit", e.text)
        self.assertIn("pending review", e.text)
        row_values = []
        for segment in rows:
            lines = e.text.splitlines()[segment["line_start"] - 1:segment["line_end"]]
            row_values.append(json.loads("\n".join(lines)))
        self.assertEqual("Declared relationship", row_values[1]["preceding_heading"])
        self.assertEqual("1:N", row_values[1]["columns"][2]["value"])
        self.assertEqual("pending review", row_values[1]["columns"][3]["value"])
        self.assertEqual("State transitions", row_values[2]["preceding_heading"])
        self.assertEqual("Model change history", row_values[3]["preceding_heading"])
        self.assertEqual("2026-09-01", row_values[3]["columns"][1]["value"])

    def test_html_merged_table_cells_do_not_claim_header_alignment(self):
        e = extract_bytes(b'<table><tr><th>A</th><th>B</th></tr><tr><td colspan="2">joined</td></tr></table>',
                          "merged.html")
        segment = next(s for s in e.metadata["segments"] if s.get("content_type") == "table_row")
        lines = e.text.splitlines()[segment["line_start"] - 1:segment["line_end"]]
        row = json.loads("\n".join(lines))
        self.assertIsNone(row["columns"][0]["header"])
        self.assertIn("html_table_spans_require_review", {i["code"] for i in e.metadata["issues"]})

    def test_html_merged_headers_suppress_following_row_alignment(self):
        raw = (b'<table><tr><th colspan="2">Amount</th><th>Status</th></tr>'
               b'<tr><td>10</td><td>USD</td><td>pending</td></tr></table>')

        result = extract_bytes(raw, "headers.html")

        lines = result.text.splitlines()
        segment = next(s for s in result.metadata["segments"] if s.get("content_type") == "table_row")
        columns = json.loads("\n".join(lines[segment["line_start"] - 1:segment["line_end"]]))["columns"]
        self.assertEqual([c["value"] for c in columns], ["10", "USD", "pending"])
        self.assertEqual([c["header"] for c in columns], [None, None, None])

    def test_html_inherited_rowspan_suppresses_alignment_without_affecting_other_tables(self):
        raw = (b'<table><tr><th>Entity</th><th>Status</th></tr>'
               b'<tr><td rowspan="2">Supplier</td><td>pending</td></tr>'
               b'<tr><td>approved</td></tr></table>'
               b'<table><tr><th>Status</th></tr><tr><td>proposed</td></tr></table>')

        result = extract_bytes(raw, "rowspan.html")

        lines = result.text.splitlines()
        rows = [json.loads("\n".join(lines[s["line_start"] - 1:s["line_end"]]))
                for s in result.metadata["segments"] if s.get("content_type") == "table_row"]
        self.assertEqual(rows[1]["columns"][0]["value"], "approved")
        self.assertIsNone(rows[1]["columns"][0]["header"])
        self.assertEqual(rows[2]["columns"][0]["header"], "Status")

    def test_html_table_cell_budget_blocks_capture(self):
        self.assert_code("extraction_budget_exceeded", extract_bytes,
                         b"<table><tr><td>one</td><td>two</td></tr></table>", "many.html",
                         limits=replace(Limits(), max_cells=1))

    def test_html_later_header_only_row_remains_citable(self):
        raw = b'<table><tr><th>Code</th></tr><tr><th>00123</th></tr></table>'
        e = extract_bytes(raw, "headers.html")
        rows = [s for s in e.metadata["segments"] if s.get("content_type") == "table_row"]
        self.assertEqual(1, len(rows))
        lines = e.text.splitlines()[rows[0]["line_start"] - 1:rows[0]["line_end"]]
        self.assertEqual("00123", json.loads("\n".join(lines))["columns"][0]["value"])
        self.assertIn("html_additional_header_row", {i["code"] for i in e.metadata["issues"]})

    def test_structured_csv_keeps_declared_headers_multiline_and_literals(self):
        raw = (ROOT / "tests/fixtures/patch_f_values.csv").read_bytes()
        e = extract_bytes(raw, "values.csv", options=Options(csv_delimiter=";", structured_csv=True,
                                                                csv_header="first-row"))
        rows = [s for s in e.metadata["segments"] if s.get("content_type") == "table_row"]
        self.assertEqual(2, len(rows))
        lines = e.text.splitlines()[rows[0]["line_start"] - 1:rows[0]["line_end"]]
        first = json.loads("\n".join(lines))
        self.assertEqual("00123", first["columns"][0]["value"])
        self.assertEqual("Amount", first["columns"][1]["header"])
        self.assertEqual("12\n.50", first["columns"][1]["value"])
        self.assertEqual("", first["columns"][3]["value"])
        self.assertIn('"=1+1"', e.text)

    def test_structured_csv_declared_header_does_not_skip_empty_first_record(self):
        self.assert_code("csv_declared_header_missing", extract_bytes, b"\nCode,Value\n00123,42\n",
                         "missing.csv", options=Options(structured_csv=True, csv_header="first-row"))

    def test_structured_csv_without_header_uses_positions_only(self):
        e = extract_bytes(b"00123,42\n", "no-header.csv", options=Options(structured_csv=True))
        segment = next(s for s in e.metadata["segments"] if s.get("content_type") == "table_row")
        lines = e.text.splitlines()[segment["line_start"] - 1:segment["line_end"]]
        row = json.loads("\n".join(lines))
        self.assertEqual("00123", row["columns"][0]["value"])
        self.assertIsNone(row["columns"][0]["header"])

    def test_structured_csv_duplicate_declared_headers_are_flagged(self):
        e = extract_bytes(b"Code,Code\n00123,42\n", "duplicate.csv",
                          options=Options(structured_csv=True, csv_header="first-row"))
        self.assertIn("csv_duplicate_declared_header", {i["code"] for i in e.metadata["issues"]})

    def test_eml_declared_legacy_charset_is_decoded(self):
        e = extract_bytes(b'Subject: Status\nContent-Type: text/plain; charset="windows-1252"\n\nCaf\xe9',
                          "encoded.eml")
        self.assertIn("Café", e.text)
        self.assertEqual("windows-1252", e.metadata["details"]["body_charsets"][0]["decoded"])

    def test_eml_bad_boundary_or_encoding_blocks(self):
        self.assert_code("eml_malformed_structure", extract_bytes,
                         b'MIME-Version: 1.0\nContent-Type: multipart/mixed; boundary="absent"\n\nBody',
                         "bad.eml")
        self.assert_code("charset_decode_error", extract_bytes,
                         b'Content-Type: text/plain; charset="utf-8"\n\n\xe9', "bad.eml")

    def test_eml_part_budget_blocks(self):
        raw = (ROOT / "tests/fixtures/patch_e_multipart.eml").read_bytes()
        self.assert_code("extraction_budget_exceeded", extract_bytes, raw, "many.eml",
                         limits=replace(Limits(), max_mime_parts=2))

    def test_eml_attached_message_is_inventory_only(self):
        raw = (b'From: outer@example.invalid\nMIME-Version: 1.0\n'
               b'Content-Type: multipart/mixed; boundary="edge"\n\n--edge\n'
               b'Content-Type: message/rfc822\nContent-Disposition: attachment; filename="forward.eml"\n\n'
               b'From: inner@example.invalid\nSubject: Secret inner subject\n\nSecret inner body\n'
               b'--edge--\n')
        e = extract_bytes(raw, "outer.eml")
        self.assertIn("forward.eml", e.text)
        self.assertNotIn("Secret inner body", e.text)
        self.assertIn("eml_attachment_not_extracted", {i["code"] for i in e.metadata["issues"]})

    def test_eml_nested_message_without_disposition_is_inventory_only(self):
        raw = (b'From: outer@example.invalid\nMIME-Version: 1.0\n'
               b'Content-Type: multipart/mixed; boundary="edge"\n\n--edge\n'
               b'Content-Type: message/rfc822\n\n'
               b'From: inner@example.invalid\nSubject: Inner subject\n\nInner body\n'
               b'--edge--\n')
        e = extract_bytes(raw, "outer.eml")
        self.assertNotIn("Inner body", e.text)
        self.assertIn("message/rfc822", e.text)

    def test_csv_basic_unicode_and_leading_zeros(self):
        e = extract_bytes("Código,Estado\n00123,Aprobado\n".encode(), "a.csv")
        self.assertIn('["00123", "Aprobado"]', e.text)
        self.assertIn("Código", e.text)
        self.assertEqual(e.metadata["status"], "extracted_needs_review")

    def test_csv_utf8_bom(self):
        self.assertNotIn("\ufeff", extract_bytes(b"\xef\xbb\xbfa,b\n1,2", "A.CSV").text)

    def test_csv_multiline_record_locator(self):
        e = extract_bytes(b'a,b\n1,"two\nlines"\n', "a.csv")
        self.assertIn("physical-lines=2-3", e.metadata["segments"][1]["locator"])
        self.assertIn('two\\nlines', e.text)

    def test_csv_delimiter_is_explicit(self):
        e = extract_bytes(b"a;b\n1;2\n", "a.csv", options=Options(csv_delimiter=";"))
        self.assertIn('["1", "2"]', e.text)
        default = extract_bytes(b"a;b\n1;2\n", "a.csv")
        self.assertIn("csv_single_column", {i["code"] for i in default.metadata["issues"]})

    def test_csv_encoding_error(self):
        self.assert_code("csv_encoding_error", extract_bytes, b"a\n\xe9", "a.csv")

    def test_csv_encoding_explicit(self):
        e = extract_bytes(b"a\n\xe9", "a.csv", options=Options(csv_encoding="cp1252"))
        self.assertIn("é", e.text)

    def test_csv_invalid_quoting_blocks(self):
        self.assert_code("csv_parse_error", extract_bytes, b'a,b\n1,"unterminated', "a.csv")

    def test_csv_formula_is_string(self):
        self.assertIn('"=1+1"', extract_bytes(b"a\n=1+1\n", "a.csv").text)

    def test_csv_ragged_warning(self):
        e = extract_bytes(b"a,b\n1\n", "a.csv")
        self.assertIn("csv_ragged_rows", {i["code"] for i in e.metadata["issues"]})

    def test_csv_cell_budget(self):
        self.assert_code("extraction_budget_exceeded", extract_bytes, b"a,b", "a.csv", limits=replace(Limits(), max_cells=1))

    def test_nul_rejected(self):
        self.assert_code("invalid_text", extract_bytes, b"a\x00", "a.csv")

    def test_docx_order_nested_table_header_footer(self):
        e = extract_bytes(docx_fixture(), "a.docx")
        self.assertLess(e.text.index("Primero"), e.text.index("Código"))
        self.assertLess(e.text.index("Código"), e.text.index("Después"))
        for phrase in ("Nested fact", "Header fact", "Footer fact", "00123"):
            self.assertIn(phrase, e.text)
        self.assertTrue(any("table/row=1/column=2" in s["locator"] for s in e.metadata["segments"]))

    def test_docx_tracked_changes_block(self):
        raw = edit_package(docx_fixture(), "word/document.xml", lambda d: d.replace(b"<w:body>", b'<w:body><w:ins w:id="1"><w:p><w:r><w:t>Unreviewed</w:t></w:r></w:p></w:ins>'))
        self.assert_code("tracked_changes_require_review", extract_bytes, raw, "a.docx")

    def test_docx_content_controls_block(self):
        raw = edit_package(docx_fixture(), "word/document.xml", lambda d: d.replace(b"<w:body>", b'<w:body><w:sdt/>'))
        self.assert_code("docx_structure_requires_export", extract_bytes, raw, "a.docx")

    def test_pptx_notes_tables_groups_hidden(self):
        e = extract_bytes(pptx_fixture(), "a.pptx")
        for phrase in ("Proveedor autorizado", "00123", "Propuesta, no decisión.", "Grouped text"):
            self.assertIn(phrase, e.text)
        self.assertEqual(e.metadata["details"]["slide_count"], 2)
        self.assertIn("hidden_slide_included", {i["code"] for i in e.metadata["issues"]})

    def test_ppt_requires_opt_in(self):
        self.assert_code("legacy_ppt_opt_in_required", extract_bytes, bytes.fromhex("D0CF11E0A1B11AE1") + b"fake", "a.ppt")

    def test_ppt_renamed_pptx_rejected(self):
        self.assert_code("format_mismatch", extract_bytes, pptx_fixture(), "a.ppt")

    def test_pdf_page_locator(self):
        e = extract_bytes(pdf_fixture(), "a.pdf")
        self.assertIn("Approved supplier 00123", e.text)
        self.assertEqual(e.metadata["segments"][0]["locator"], "PDF/page=1")

    def test_pdf_mixed_text_and_blank_visible(self):
        e = extract_bytes(pdf_fixture(mixed=True), "a.pdf")
        self.assertIn("pdf_page_without_text", {i["code"] for i in e.metadata["issues"]})
        self.assertIn("pdf_page_without_text", e.text)

    def test_pdf_empty_text_blocks(self):
        self.assert_code("no_extractable_text", extract_bytes, pdf_fixture(blank=True), "a.pdf")

    def test_pdf_encrypted_blocks(self):
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter()
        writer.add_page(PdfReader(BytesIO(pdf_fixture())).pages[0])
        writer.encrypt("test-fixture-password")
        stream = BytesIO()
        writer.write(stream)
        self.assert_code("encrypted_document", extract_bytes, stream.getvalue(), "a.pdf")

    def test_pdf_page_budget(self):
        self.assert_code("extraction_budget_exceeded", extract_bytes, pdf_fixture(mixed=True), "a.pdf", limits=replace(Limits(), max_pages_or_slides=1))

    def test_xlsx_cells_formulas_and_caches_separate(self):
        e = extract_bytes(xlsx_fixture(), "a.xlsx")
        self.assertIn('"formula": "SUM(B2,1)"', e.text)
        self.assertIn('"cached_value": "13.50"', e.text)
        self.assertIn('"cached_value_present": false', e.text)
        self.assertIn('"value": "00123"', e.text)
        self.assertIn('"value": "12.50"', e.text)
        self.assertIn('"column_hidden": true', e.text)
        self.assertIn('"row_hidden": true', e.text)
        self.assertEqual(e.metadata["details"]["excel_date_system"], "1904")
        self.assertIn("hidden_sheet_included", {i["code"] for i in e.metadata["issues"]})
        self.assertIn("xlsx_merged_cells", {i["code"] for i in e.metadata["issues"]})

    def test_xlsx_large_sparse_coordinate_does_not_expand_grid(self):
        raw = edit_package(xlsx_fixture(), "xl/worksheets/sheet1.xml", lambda d: d.replace(b'r="F2"', b'r="XFD1048576"'))
        e = extract_bytes(raw, "a.xlsx")
        self.assertIn("cell=XFD1048576", e.text)
        self.assertEqual(len(e.metadata["segments"]), 7)

    def test_xlsx_style_only_cells_not_source_values(self):
        raw = edit_package(xlsx_fixture(), "xl/worksheets/sheet1.xml", lambda d: d.replace(b"</sheetData>", b'<row r="4"><c r="A4" s="1"/></row></sheetData>'))
        self.assertNotIn("cell=A4", extract_bytes(raw, "a.xlsx").text)

    def test_xlsx_external_relationship_warning(self):
        raw = edit_package(xlsx_fixture(), "xl/_rels/workbook.xml.rels", lambda d: d.replace(b"</Relationships>", b'<Relationship Id="e" Target="https://invalid.example/" TargetMode="External"/></Relationships>'))
        e = extract_bytes(raw, "a.xlsx")
        self.assertIn("external_relationships_not_fetched", {i["code"] for i in e.metadata["issues"]})

    def test_xlsx_duplicate_cell_coordinates_block(self):
        raw = edit_package(xlsx_fixture(), "xl/worksheets/sheet1.xml", lambda d: d.replace(b'r="B2"', b'r="A2"'))
        self.assert_code("invalid_container", extract_bytes, raw, "a.xlsx")

    def test_xlsx_duplicate_sheet_names_block(self):
        raw = edit_package(xlsx_fixture(), "xl/workbook.xml", lambda d: d.replace(b'</sheets>', b'<sheet name="Suppliers" sheetId="2" r:id="rId1"/></sheets>'))
        self.assert_code("invalid_container", extract_bytes, raw, "a.xlsx")

    def test_unicode_line_separators_do_not_shift_citations(self):
        e = extract_bytes("header\nvalue\u2028continuation\n".encode(), "a.csv")
        self.assertIn("\\u2028", e.text)
        self.assertNotIn("\u2028", e.text)
        lines = e.text.splitlines()
        last = e.metadata["segments"][-1]
        self.assertIn("value", lines[last["line_start"] - 1])

    def test_xlsx_cell_budget(self):
        self.assert_code("extraction_budget_exceeded", extract_bytes, xlsx_fixture(), "a.xlsx", limits=replace(Limits(), max_cells=1))

    def test_zip_member_budget(self):
        self.assert_code("archive_budget_exceeded", extract_bytes, xlsx_fixture(), "a.xlsx", limits=replace(Limits(), max_zip_members=1))

    def test_zip_traversal_blocks(self):
        raw = edit_package(xlsx_fixture(), "../bad.xml", lambda _: b"<bad/>")
        self.assert_code("invalid_container", extract_bytes, raw, "a.xlsx")

    def test_xml_dtd_blocks(self):
        raw = edit_package(xlsx_fixture(), "bad.xml", lambda _: b'<!DOCTYPE x [<!ENTITY t "text">]><x>&t;</x>')
        self.assert_code("unsafe_or_invalid_xml", extract_bytes, raw, "a.xlsx")

    def test_macro_package_blocks(self):
        raw = edit_package(xlsx_fixture(), "xl/vbaProject.bin", lambda _: b"fake")
        self.assert_code("active_content_unsupported", extract_bytes, raw, "a.xlsx")

    def test_wrong_container_format(self):
        self.assert_code("format_mismatch", extract_bytes, xlsx_fixture(), "a.docx")

    def test_invalid_binary_file_does_not_succeed(self):
        self.assert_code("invalid_container", extract_bytes, b"Not a document", "a.docx")

    def test_unsupported_extension(self):
        self.assert_code("unsupported_extension", extract_bytes, b"content", "a.xls")

    def test_input_budget(self):
        self.assert_code("input_budget_exceeded", extract_bytes, b"abcd", "a.csv", limits=replace(Limits(), max_input_bytes=2))

    def test_output_budget_no_truncation(self):
        self.assert_code("extraction_budget_exceeded", extract_bytes, b"abc", "a.csv", limits=replace(Limits(), max_output_chars=10))

    def test_empty_document_blocks(self):
        self.assert_code("empty_file", extract_bytes, b"", "a.csv")

    def test_dependency_missing_is_explicit(self):
        with patch("kb_document_extractors.load_dependency", side_effect=ExtractionError("dependency_missing", "synthetic")):
            self.assert_code("dependency_missing", extract_bytes, pdf_fixture(), "a.pdf")

    def test_hashes_locators_repeatability(self):
        for name, raw in (("a.csv", b"a,b\n1,2\n"), ("a.xlsx", xlsx_fixture()), ("a.docx", docx_fixture()), ("a.pptx", pptx_fixture()), ("a.pdf", pdf_fixture())):
            with self.subTest(name=name):
                e = extract_bytes(raw, name)
                self.assertEqual(e, extract_bytes(raw, name))
                self.assertEqual(digest(raw), e.metadata["source_sha256"])
                self.assertEqual(digest(e.text.encode()), e.metadata["text_sha256"])
                previous = 0
                lines = e.text.splitlines()
                for s in e.metadata["segments"]:
                    self.assertGreater(s["line_start"], previous)
                    self.assertLessEqual(s["line_end"], len(lines))
                    self.assertTrue("\n".join(lines[s["line_start"]-1:s["line_end"]]).strip())
                    previous = s["line_end"]


class SnapshotTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / "input.csv"
        self.source.write_bytes(b"code,status\n00123,approved\n")
        self.out = self.root / "snapshot"

    def cli(self, *args):
        return subprocess.run([sys.executable, str(ROOT / "tools/kb_extract_document.py"), *map(str, args)], capture_output=True, text=True, timeout=20)

    def capture(self):
        result = self.cli("--source", self.source, "--out", self.out)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_cli_roundtrip_and_source_unchanged(self):
        raw = self.source.read_bytes()
        self.capture()
        self.assertEqual(self.source.read_bytes(), raw)
        self.assertEqual(verify(self.out, self.source)["status"], "integrity_passed")
        result = self.cli("--verify", self.out, "--compare-source", self.source)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_existing_output_not_overwritten(self):
        self.capture()
        before = (self.out / "extraction.json").read_bytes()
        result = self.cli("--source", self.source, "--out", self.out)
        self.assertEqual(result.returncode, 2)
        self.assertIn("output_exists", result.stderr)
        self.assertEqual(before, (self.out / "extraction.json").read_bytes())

    def test_raw_tampering_rejected(self):
        self.capture()
        (self.out / "original.csv").write_bytes(b"tampered")
        with self.assertRaisesRegex(ExtractionError, "original_hash_mismatch"):
            verify(self.out)

    def test_text_tampering_rejected(self):
        self.capture()
        (self.out / "evidence.md").write_text("tampered")
        with self.assertRaisesRegex(ExtractionError, "text_hash_mismatch"):
            verify(self.out)

    def test_metadata_tampering_rejected(self):
        self.capture()
        (self.out / "extraction.json").write_text("{}")
        with self.assertRaisesRegex(ExtractionError, "metadata_hash_mismatch"):
            verify(self.out)

    def test_source_change_is_detected(self):
        self.capture()
        self.source.write_bytes(b"new source")
        with self.assertRaisesRegex(ExtractionError, "source_changed_since_capture"):
            verify(self.out, self.source)

    def test_file_symlink_rejected(self):
        link = self.root / "link.csv"
        link.symlink_to(self.source)
        with self.assertRaisesRegex(ExtractionError, "symlink_not_allowed"):
            read_stable(link, 100)

    def test_failed_parse_does_not_create_snapshot(self):
        self.source.write_bytes(b"a\x00")
        result = self.cli("--source", self.source, "--out", self.out)
        self.assertEqual(result.returncode, 2)
        self.assertFalse(self.out.exists())


@unittest.skipUnless(os.environ.get("KB_TEST_LEGACY_PPT") == "1", "Set KB_TEST_LEGACY_PPT=1 to exercise installed LibreOffice")
class RealLegacyPptTests(unittest.TestCase):
    def test_real_legacy_ppt_conversion(self):
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        self.assertIsNotNone(executable, "LibreOffice is required for the requested conversion test")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory).resolve()
            source = root / "fixture.pptx"
            source.write_bytes(pptx_fixture())
            command = [executable, "-env:UserInstallation=" + (root / "profile").as_uri(), "--headless", "--convert-to", "ppt:MS PowerPoint 97", "--outdir", str(root), str(source)]
            result = subprocess.run(command, capture_output=True, timeout=45)
            self.assertEqual(result.returncode, 0, result.stderr)
            raw = (root / "fixture.ppt").read_bytes()
            e = extract_bytes(raw, "fixture.ppt", options=Options(allow_legacy_ppt=True, soffice=executable))
            self.assertIn("Proveedor autorizado", e.text)
            self.assertIn("00123", e.text)
            self.assertIn("LibreOffice", e.metadata["parsers"])
            self.assertIn("legacy_ppt_conversion_requires_review", {i["code"] for i in e.metadata["issues"]})
            self.assertEqual(digest(raw), e.metadata["source_sha256"])


if __name__ == "__main__":
    unittest.main()
