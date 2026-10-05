"""Synthetic checks for versioned quotable representations and segment bindings."""
import json
import shutil
import struct
import sys
import tempfile
import unittest
import zlib
from io import BytesIO
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

from docx import Document

import kb_inventory
from kb_check import check_records, check_run, load_manifest
from kb_common import KBError, read_json, sha, utc_now
from kb_inventory import inventory
from kb_packet import build_packets
from kb_publish import prepare_review, publish
from kb_reanchor import reanchor


def synthetic_png() -> bytes:
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
            + chunk(b"IEND", b""))


class SegmentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        source = self.root / "source"
        source.mkdir()
        doc = Document()
        doc.add_paragraph("Alpha requirement.")
        doc.add_picture(BytesIO(synthetic_png()))
        buffer = BytesIO()
        doc.save(buffer)
        (source / "source.docx").write_bytes(buffer.getvalue())
        project = self.root / "project"
        (project / "config").mkdir(parents=True)
        self.config = project / "config" / "project.json"
        self.config.write_text(json.dumps({
            "schema_version": "0.1",
            "project": {"id": "example", "name": "Example", "description": "Synthetic"},
            "sources": [{"id": "notes", "type": "requirements", "path": str(source),
                         "include": ["**/*.docx"], "exclude": []}],
            "global_exclude": [],
            "knowledge": {"path": "knowledge", "approved": "knowledge/approved"},
            "runs": {"path": "runs"},
        }), encoding="utf-8")
        self.source = source

    def records_03(self, run, manifest):
        f = manifest["files"][0]
        segment_data = read_json(run / "segments.snapshot.json")
        segment = next(s for s in segment_data["segments"] if s["content_type"] == "unclassified")
        lines = (run / f["snapshot_path"]).read_text().splitlines()
        citation = {"evidence_id": f["evidence_id"], "segment_id": segment["segment_id"],
                    "representation_sha256": segment["representation_sha256"],
                    "start_line": segment["line_start"], "end_line": segment["line_start"],
                    "quote": lines[segment["line_start"] - 1]}
        data = {"schema_version": "0.3", "project_id": "example", "run_id": manifest["run_id"],
                "coverage": [{"evidence_id": f["evidence_id"], "disposition": "used",
                              "note": "Synthetic citation; not a human approval."}],
                "segment_coverage": [{"segment_id": s["segment_id"],
                                      "disposition": "used" if s["segment_id"] == segment["segment_id"]
                                      else "triaged_out",
                                      "note": "Synthetic segment coverage; no visual claim.",
                                      **({} if s["segment_id"] == segment["segment_id"] else
                                         {"method": "Exclude embedded asset metadata from the synthetic text claim."})}
                                     for s in segment_data["segments"]],
                "records": [{"id": "REQ-001", "kind": "requirement", "title": "Synthetic text",
                             "statement": "The synthetic document contains Alpha requirement.",
                             "epistemic_status": "observed", "evidence": [citation],
                             "relations": [], "open_questions": [], "investigation": None}]}
        return data

    def save_records(self, path, data):
        path.write_text(json.dumps(data), encoding="utf-8")
        return path

    def test_segment_triage_is_reported_and_requires_explicit_review(self):
        run, manifest = inventory(self.config, "triage", documents=True)
        records = self.records_03(run, manifest)
        self.save_records(run / "proposals/records.json", records)
        review = prepare_review(self.config, run)
        self.assertTrue(review["triaged_segment_ids"])
        report = (run / "review-report.md").read_text()
        for segment_id in review["triaged_segment_ids"]:
            self.assertIn(segment_id, report)
        review.update(decision="approve", reviewer="Synthetic test operator (simulated)",
                      reviewed_at=utc_now(), notes="UNIT TEST ONLY. Not a real human approval.")
        review["checks"] = dict.fromkeys(review["checks"], True)
        review["acknowledged_warnings"] = dict(review["warning_summary"]["counts"])
        review_path = run / "review.json"
        self.save_records(review_path, review)
        with self.assertRaisesRegex(KBError, "segment_triage_acknowledged"):
            publish(self.config, run, review_path, True)
        review["segment_triage_acknowledged"] = True
        self.save_records(review_path, review)
        self.assertTrue((publish(self.config, run, review_path, True) / "review-report.md").exists())

    def test_changed_derived_text_changes_document_evidence_identity(self):
        first, old = inventory(self.config, "first", documents=True)
        original_extract = kb_inventory.extract_document_snapshot

        def changed_extract(*args, **kwargs):
            metadata, text = original_extract(*args, **kwargs)
            text = text.replace("Alpha requirement.", "Bravo requirement.")
            metadata["text_sha256"] = sha(text.encode("utf-8"))
            metadata["text_bytes"] = len(text.encode("utf-8"))
            return metadata, text

        with mock.patch.object(kb_inventory, "extract_document_snapshot", side_effect=changed_extract):
            second, new = inventory(self.config, "second", baseline=first, documents=True)
        before, after = old["files"][0], new["files"][0]
        self.assertEqual(before["sha256"], after["sha256"])
        self.assertNotEqual(before["document"]["text_sha256"], after["document"]["text_sha256"])
        self.assertNotEqual(before["evidence_id"], after["evidence_id"])
        self.assertIn(before["logical_id"], new["delta"]["modified"])
        self.assertEqual(check_run(first)["status"], "passed")
        self.assertEqual(check_run(second)["status"], "passed")

    def test_parser_version_changes_representation_identity_even_if_text_matches(self):
        first, old = inventory(self.config, "first", documents=True)
        self.save_records(first / "proposals" / "records.json", self.records_03(first, old))
        original_extract = kb_inventory.extract_document_snapshot

        def new_parser(*args, **kwargs):
            metadata, text = original_extract(*args, **kwargs)
            metadata["parsers"]["synthetic-parser-version"] = "next"
            return metadata, text

        with mock.patch.object(kb_inventory, "extract_document_snapshot", side_effect=new_parser):
            second, new = inventory(self.config, "second", baseline=first, documents=True)
        self.assertEqual(old["files"][0]["sha256"], new["files"][0]["sha256"])
        self.assertEqual(old["files"][0]["document"]["text_sha256"],
                         new["files"][0]["document"]["text_sha256"])
        self.assertNotEqual(old["files"][0]["evidence_id"], new["files"][0]["evidence_id"])
        self.assertIn(old["files"][0]["logical_id"], new["delta"]["modified"])
        self.assertEqual(check_run(second)["status"], "passed")
        moved, report = reanchor(first, second)
        self.assertEqual(report["citation_status_counts"], {"representation_changed": 1})
        self.assertNotEqual(moved["records"][0]["evidence"][0]["segment_id"],
                            next(s for s in read_json(first / "segments.snapshot.json")["segments"]
                                 if s["content_type"] == "unclassified")["segment_id"])

    def test_segment_inventory_binds_locators_and_reports_unavailable_visuals(self):
        run, manifest = inventory(self.config, "first", documents=True)
        data = read_json(run / "segments.snapshot.json")
        segment = next(s for s in data["segments"] if s["content_type"] == "unclassified")
        self.assertEqual(segment["evidence_id"], manifest["files"][0]["evidence_id"])
        self.assertEqual(segment["content_type"], "unclassified")
        self.assertEqual(segment["extraction_status"], "extracted_needs_review")
        self.assertTrue(segment["original_locator"].startswith("DOCX/"))
        self.assertTrue(any(i["code"] == "docx_embedded_image_not_extracted" and i["status"] == "unavailable"
                            for i in data["issues"]))
        self.assertIn("embedded_asset", {s["content_type"] for s in data["segments"]})
        index = build_packets(run)
        self.assertIn(segment["segment_id"], index["packets"][0]["segment_ids"])
        self.assertIn(segment["original_locator"], {s["original_locator"] for s in index["segments"]})

    def test_segment_inventory_rejects_tampering_even_after_resealing_manifest(self):
        run, _ = inventory(self.config, "first", documents=True)
        copy = self.root / "tampered"
        shutil.copytree(run, copy)
        snapshot = copy / "segments.snapshot.json"
        snapshot.write_bytes(snapshot.read_bytes() + b" ")
        with self.assertRaisesRegex(KBError, "Segment inventory checksum"):
            load_manifest(copy)
        shutil.rmtree(copy)
        shutil.copytree(run, copy)
        snapshot = copy / "segments.snapshot.json"
        data = read_json(snapshot)
        data["segments"][0]["original_locator"] = "invented/locator"
        snapshot.write_text(json.dumps(data), encoding="utf-8")
        manifest = read_json(copy / "manifest.json")
        manifest["segment_inventory"]["sha256"] = sha(snapshot.read_bytes())
        raw = json.dumps(manifest).encode("utf-8")
        (copy / "manifest.json").write_bytes(raw)
        (copy / "manifest.sha256").write_text(sha(raw) + "\n")
        with self.assertRaisesRegex(KBError, "Segment inventory disagrees"):
            load_manifest(copy)

    def test_records_03_require_exact_segment_and_complete_segment_coverage(self):
        run, manifest = inventory(self.config, "first", documents=True)
        good = self.records_03(run, manifest)
        path = self.save_records(self.root / "good.json", good)
        self.assertEqual(check_records(run, manifest, path, stage2=True)["segment_coverage_counts"],
                         {"used": 1, "triaged_out": len(good["segment_coverage"]) - 1})
        bad = json.loads(json.dumps(good))
        bad["records"][0]["evidence"][0]["representation_sha256"] = "0" * 64
        with self.assertRaisesRegex(KBError, "bind one frozen segment"):
            check_records(run, manifest, self.save_records(self.root / "bad-rep.json", bad))
        bad = json.loads(json.dumps(good))
        bad["segment_coverage"] = []
        with self.assertRaisesRegex(KBError, "every frozen segment"):
            check_records(run, manifest, self.save_records(self.root / "bad-coverage.json", bad))
        bad = json.loads(json.dumps(good))
        bad["records"][0]["evidence"][0].update(start_line=1, end_line=1,
                                                   quote="# Derived document evidence")
        with self.assertRaisesRegex(KBError, "bind one frozen segment"):
            check_records(run, manifest, self.save_records(self.root / "bad-range.json", bad))

    def test_reanchor_segment_citation_after_document_lines_move(self):
        first, old = inventory(self.config, "first", documents=True)
        records = self.records_03(first, old)
        self.save_records(first / "proposals" / "records.json", records)
        doc = Document()
        doc.add_paragraph("Intro context.")
        doc.add_paragraph("Additional intro context.")
        doc.add_paragraph("Alpha requirement.")
        buffer = BytesIO()
        doc.save(buffer)
        (self.source / "source.docx").write_bytes(buffer.getvalue())
        second, new = inventory(self.config, "second", baseline=first, documents=True)
        moved, report = reanchor(first, second)
        self.assertEqual(moved["schema_version"], "0.3")
        self.assertEqual(report["citation_status_counts"], {"moved": 1})
        ev = moved["records"][0]["evidence"][0]
        segment = next(s for s in read_json(second / "segments.snapshot.json")["segments"]
                       if s["segment_id"] == ev["segment_id"])
        self.assertEqual(segment["representation_sha256"], ev["representation_sha256"])
        self.assertEqual(check_records(second, new, second / "work" / "reanchored-records.json")["record_count"], 1)

    def test_reanchor_drops_quote_when_extraction_no_longer_supports_it(self):
        first, old = inventory(self.config, "first", documents=True)
        self.save_records(first / "proposals" / "records.json", self.records_03(first, old))
        doc = Document()
        doc.add_paragraph("Different material only.")
        buffer = BytesIO()
        doc.save(buffer)
        (self.source / "source.docx").write_bytes(buffer.getvalue())
        second, _ = inventory(self.config, "second", baseline=first, documents=True)
        moved, report = reanchor(first, second)
        self.assertEqual(moved["records"], [])
        self.assertEqual(report["citation_status_counts"], {"vanished": 1})
        self.assertEqual(report["dropped_record_ids"], ["REQ-001"])

    def test_segment_citation_can_span_packet_boundaries(self):
        (self.source / "long.txt").write_text("".join(f"line {i:03d} context\n" for i in range(100)),
                                              encoding="utf-8")
        cfg = read_json(self.config)
        cfg["sources"][0]["include"] = ["**/*"]
        self.config.write_text(json.dumps(cfg), encoding="utf-8")
        run, manifest = inventory(self.config, "first", documents=True)
        text_file = next(f for f in manifest["files"] if f["relative_path"] == "long.txt")
        segments = read_json(run / "segments.snapshot.json")["segments"]
        text_segment = next(s for s in segments if s["evidence_id"] == text_file["evidence_id"])
        index = build_packets(run, max_chars=1800)
        packet_rows = [p for p in index["packets"] if p["evidence_id"] == text_file["evidence_id"]]
        self.assertGreater(len(packet_rows), 1)
        self.assertIn(text_segment["segment_id"], packet_rows[0]["segment_ids"])
        self.assertIn(text_segment["segment_id"], packet_rows[1]["segment_ids"])
        start, end = packet_rows[0]["end_line"], packet_rows[1]["start_line"]
        lines = (run / text_file["snapshot_path"]).read_text().splitlines()
        data = {"schema_version": "0.3", "project_id": "example", "run_id": manifest["run_id"],
                "coverage": [{"evidence_id": f["evidence_id"],
                              "disposition": "used" if f is text_file else "reviewed_no_record",
                              "note": "Synthetic packet-boundary check."} for f in manifest["files"]],
                "segment_coverage": [{"segment_id": s["segment_id"],
                                      "disposition": "used" if s["segment_id"] == text_segment["segment_id"]
                                      else "reviewed_no_record", "note": "Synthetic check."}
                                     for s in segments],
                "records": [{"id": "REQ-001", "kind": "requirement", "title": "Boundary quote",
                             "statement": "Two adjacent synthetic lines are quoted.",
                             "epistemic_status": "observed",
                             "evidence": [{"evidence_id": text_file["evidence_id"],
                                           "segment_id": text_segment["segment_id"],
                                           "representation_sha256": text_segment["representation_sha256"],
                                           "start_line": start, "end_line": end,
                                           "quote": "\n".join(lines[start - 1:end])}],
                             "relations": [], "open_questions": [], "investigation": None}]}
        self.assertEqual(check_records(run, manifest, self.save_records(self.root / "boundary.json", data),
                                       stage2=True)["citation_count"], 1)


if __name__ == "__main__":
    unittest.main()
