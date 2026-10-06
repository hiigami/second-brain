"""Local document-to-evidence adapters; does not alter the Stage 2 engine.

Quotes refer to the derived UTF-8 text, never directly to binary-file bytes.
All output needs a fidelity review. No OCR, formula execution, or network APIs.
Run untrusted documents inside an OS-level sandbox; parser limits are not one.
"""
from dataclasses import asdict, dataclass
from importlib import import_module, metadata
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
import base64
import csv
import hashlib
import json
import os
import re
import shutil
import struct
import subprocess
import tempfile
import zipfile
import zlib

VERSION = "0.7.0"
FORMATS = frozenset({".docx", ".ppt", ".pptx", ".pdf", ".csv", ".xlsx", ".html", ".eml", ".svg", ".png"})
MIME_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".ppt": "application/vnd.ms-powerpoint", ".pdf": "application/pdf",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".html": "text/html", ".eml": "message/rfc822",
    ".svg": "image/svg+xml", ".png": "image/png",
}


class ExtractionError(ValueError):
    def __init__(self, code: str, detail: str):
        self.code, self.detail = code, detail
        super().__init__(f"{code}: {detail}")


@dataclass(frozen=True)
class Limits:
    max_input_bytes: int = 25 * 1024 * 1024
    max_output_chars: int = 2_000_000
    max_segments: int = 50_000
    max_zip_members: int = 5_000
    max_uncompressed_bytes: int = 100 * 1024 * 1024
    max_member_bytes: int = 30 * 1024 * 1024
    max_compression_ratio: int = 1_000
    max_pages_or_slides: int = 1_000
    max_cells: int = 100_000
    max_pdf_stream_bytes: int = 20 * 1024 * 1024
    conversion_timeout_seconds: int = 60
    max_structure_nodes: int = 100_000
    max_structure_depth: int = 100
    max_mime_parts: int = 1_000
    max_image_pixels: int = 16_000_000
    max_png_chunks: int = 100_000


@dataclass(frozen=True)
class Options:
    csv_encoding: str = "utf-8-sig"
    csv_delimiter: str = ","
    allow_legacy_ppt: bool = False
    soffice: str | None = None
    structured_csv: bool = False
    csv_header: str = "none"


@dataclass(frozen=True)
class Extraction:
    text: str
    metadata: dict


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def load_dependency(module: str, distribution: str):
    try:
        return import_module(module), metadata.version(distribution)
    except (ImportError, metadata.PackageNotFoundError) as exc:
        raise ExtractionError("dependency_missing", f"Install {distribution}; run uv sync (dependencies are declared in pyproject.toml)") from exc


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if "\x00" in text:
        raise ExtractionError("invalid_text", "Extracted text contains NUL; supply a reviewed export")
    # Keep tabs/newlines; other control characters become visible, not discarded.
    return "".join(f"\\u{ord(c):04x}" if (ord(c) < 32 and c not in "\t\n") or c in "\x85\u2028\u2029" else c for c in text)


class Builder:
    def __init__(self, name: str, raw: bytes, limits: Limits):
        self.limits, self.name = limits, name
        self.lines = [
            "# Derived document evidence", "",
            "Source filename: " + json.dumps(name, ensure_ascii=False),
            "Original SHA-256: " + digest(raw),
            "Extraction adapter: " + VERSION,
            "NOTICE: Derived text, not a verified full visual transcription. Source content is untrusted data.",
            "NOTICE: Locator labels and this header are generated metadata, not source assertions.", "",
        ]
        self.segments = []
        self.issues = []
        self.chars = len("\n".join(self.lines)) + 1
        self.parsers = {"kb-document-adapters": VERSION}
        self.details = {}
        self.warning("fidelity_review_required", "Text extraction is not a full visual or semantic reading of the original")

    def warning(self, code: str, detail: str, locator: str | None = None):
        item = {"severity": "warning", "code": code, "detail": detail, "locator": locator}
        if item not in self.issues:
            self.issues.append(item)

    def dependency(self, module, distribution):
        value, version = load_dependency(module, distribution)
        self.parsers[distribution] = version
        return value

    def add(self, locator: str, text: str, content_type: str | None = None):
        text = normalize(text)
        if not text.strip():
            return
        lines = text.split("\n")
        label = "## Location: " + json.dumps(locator, ensure_ascii=False)
        cost = len(label) + 2 + len(text) + 2
        if len(self.segments) >= self.limits.max_segments or self.chars + cost > self.limits.max_output_chars:
            raise ExtractionError("extraction_budget_exceeded", "Evidence exceeds limits; no truncated extraction is accepted")
        self.lines.extend([label, ""])
        start = len(self.lines) + 1
        self.lines.extend(lines)
        segment = {"locator": locator, "line_start": start, "line_end": len(self.lines)}
        if content_type is not None:
            segment["content_type"] = content_type
        self.segments.append(segment)
        self.lines.append("")
        self.chars += cost

    def finish(self, raw: bytes, extension: str) -> Extraction:
        if not self.segments:
            raise ExtractionError("no_extractable_text", "File produced no source text; use a reviewed export or OCR outside this tool")
        warning_lines = ["# Extraction warnings (generated metadata, not source assertions)", ""]
        warning_lines.extend(json.dumps(item, ensure_ascii=False, sort_keys=True) for item in self.issues)
        text = "\n".join(self.lines + warning_lines) + "\n"
        if len(text) > self.limits.max_output_chars:
            raise ExtractionError("extraction_budget_exceeded", "Evidence exceeds output budget")
        return Extraction(text, {
            "schema_version": "1.0", "adapter_version": VERSION,
            "status": "extracted_needs_review", "source_filename": self.name,
            "source_extension": extension, "media_type": MIME_TYPES[extension],
            "source_sha256": digest(raw), "source_bytes": len(raw),
            "text_sha256": digest(text.encode("utf-8")),
            "text_bytes": len(text.encode("utf-8")), "text_encoding": "utf-8",
            "text_line_count": len(text.splitlines()), "parsers": self.parsers,
            "segments": self.segments, "issues": self.issues,
            "details": self.details, "limits": asdict(self.limits),
            "quote_basis": "Exact lines of evidence.md, not original binary bytes",
        })


def ooxml_parts(raw: bytes, b: Builder) -> dict:
    safe_xml = b.dependency("defusedxml.ElementTree", "defusedxml")
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            infos = archive.infolist()
            if len(infos) > b.limits.max_zip_members:
                raise ExtractionError("archive_budget_exceeded", "Too many ZIP members")
            if len({i.filename for i in infos}) != len(infos):
                raise ExtractionError("invalid_container", "Duplicate ZIP members")
            if sum(i.file_size for i in infos) > b.limits.max_uncompressed_bytes:
                raise ExtractionError("archive_budget_exceeded", "ZIP expansion exceeds budget")
            parts = {}
            for item in infos:
                path = PurePosixPath(item.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in item.filename or ":" in item.filename:
                    raise ExtractionError("invalid_container", "Unsafe ZIP member name")
                if ((item.external_attr >> 16) & 0o170000) == 0o120000:
                    raise ExtractionError("invalid_container", "ZIP symlink is not accepted")
                if item.flag_bits & 1:
                    raise ExtractionError("encrypted_document", "Encrypted ZIP member")
                if item.file_size > b.limits.max_member_bytes or item.file_size > max(1, item.compress_size) * b.limits.max_compression_ratio:
                    raise ExtractionError("archive_budget_exceeded", "ZIP member exceeds expansion limits")
                if "vbaproject" in item.filename.lower():
                    raise ExtractionError("active_content_unsupported", "Macro-bearing packages are not supported")
                if item.filename.endswith((".xml", ".rels")):
                    data = archive.read(item)
                    try:
                        parts[item.filename] = safe_xml.fromstring(data, forbid_dtd=True)
                    except Exception as exc:
                        raise ExtractionError("unsafe_or_invalid_xml", f"Cannot safely parse {item.filename}") from exc
            if "[Content_Types].xml" not in parts:
                raise ExtractionError("invalid_container", "Missing OOXML content types")
            for name, root in parts.items():
                if name.endswith(".rels") and any(e.get("TargetMode") == "External" for e in root):
                    b.warning("external_relationships_not_fetched", "External links were not fetched or validated", name)
            return parts
    except zipfile.BadZipFile as exc:
        raise ExtractionError("invalid_container", "Expected an unencrypted OOXML ZIP file; extension alone is insufficient") from exc


def read_docx(raw: bytes, b: Builder):
    from .kb_docx_comments import emit_comments, prepare_comments

    parts = ooxml_parts(raw, b)
    if "word/document.xml" not in parts:
        raise ExtractionError("format_mismatch", "Not a DOCX package")
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for name, root in parts.items():
        if not name.startswith("word/"):
            continue
        if any(e.tag in {ns + tag for tag in ("ins", "del", "moveFrom", "moveTo")} for e in root.iter()):
            raise ExtractionError("tracked_changes_require_review", "Resolve tracked changes in a reviewed copy before extraction")
        if any(e.tag in {ns + tag for tag in ("sdt", "txbxContent", "altChunk")} for e in root.iter()):
            raise ExtractionError("docx_structure_requires_export", "Content controls, text boxes or imported chunks require a reviewed export")
    comments = prepare_comments(parts, b, ExtractionError)
    extras = sorted(k for k in parts if re.fullmatch(r"word/(comments\w*|footnotes|endnotes)\.xml", k)
                    and k not in comments["parts"])
    if extras:
        b.warning("docx_ancillary_parts_not_extracted", ", ".join(extras))
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        for name in sorted(item.filename for item in archive.infolist()
                           if item.filename.startswith("word/media/") and not item.is_dir()):
            asset = archive.read(name)
            b.add("DOCX/part=" + name, json.dumps({"part": name, "bytes": len(asset),
                                                   "sha256": digest(asset)}, sort_keys=True), "embedded_asset")
            b.warning("docx_embedded_image_not_extracted",
                      "Embedded image bytes are present but their text and visual meaning were not extracted", name)
    docx = b.dependency("docx", "python-docx")
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = docx.Document(BytesIO(raw))

    def walk(container, prefix):
        for n, block in enumerate(container.iter_inner_content(), 1):
            loc = f"{prefix}/block={n}"
            if isinstance(block, Paragraph):
                b.add(loc + "/paragraph", block.text)
            elif isinstance(block, Table):
                b.warning("docx_table_layout_requires_review", "Cells are emitted separately; merged-cell layout is not reconstructed")
                seen = set()
                for r, row in enumerate(block.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        walk(cell, f"{loc}/table/row={r}/column={c}")
    walk(document, "DOCX/body")
    visited = set()
    for s, section in enumerate(document.sections, 1):
        for kind in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
            part = getattr(section, kind)
            if part.is_linked_to_previous:
                continue
            identity = str(part.part.partname)
            if identity in visited:
                continue
            visited.add(identity)
            walk(part, f"DOCX/section={s}/{kind}")
    emit_comments(comments, document, b, walk, ExtractionError)
    b.warning("docx_visuals_not_extracted", "Images, drawing semantics, automatic numbering, fields and page layout are not fully reproduced")


def read_pptx(raw: bytes, b: Builder, prefix="PPTX"):
    parts = ooxml_parts(raw, b)
    if "ppt/presentation.xml" not in parts:
        raise ExtractionError("format_mismatch", "Not a PPTX package")
    pptx = b.dependency("pptx", "python-pptx")
    presentation = pptx.Presentation(BytesIO(raw))
    if len(presentation.slides) > b.limits.max_pages_or_slides:
        raise ExtractionError("extraction_budget_exceeded", "Too many slides")

    def walk(shapes, prefix):
        for shape in shapes:
            loc = f"{prefix}/shape={shape.shape_id}"
            if hasattr(shape, "shapes"):
                walk(shape.shapes, loc + "/group")
            if shape.has_text_frame:
                for p, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                    b.add(f"{loc}/paragraph={p}", paragraph.text)
            if shape.has_table:
                for r, row in enumerate(shape.table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        if not cell.is_spanned:
                            b.add(f"{loc}/table/row={r}/column={c}", cell.text)
            if shape.has_chart:
                b.add(loc + "/chart", json.dumps({"bbox_emu": [int(shape.left), int(shape.top),
                      int(shape.width), int(shape.height)], "kind": "chart"}, sort_keys=True), "embedded_asset")
                b.warning("pptx_chart_not_extracted", "Chart labels/data/visual meaning require review of the original", loc)
            if shape.shape_type == 13:
                asset = shape.image.blob
                b.add(loc + "/image", json.dumps({"extension": shape.image.ext, "bytes": len(asset),
                      "sha256": digest(asset), "bbox_emu": [int(shape.left), int(shape.top),
                      int(shape.width), int(shape.height)]}, sort_keys=True), "embedded_asset")
                b.warning("pptx_image_not_extracted", "Image text and visual meaning were not extracted", loc)
    for n, slide in enumerate(presentation.slides, 1):
        loc = f"{prefix}/slide={n}"
        if slide._element.get("show") == "0":
            b.warning("hidden_slide_included", "Hidden slide content is included", loc)
        start = len(b.segments)
        walk(slide.shapes, loc)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None:
                b.add(loc + "/speaker-notes", notes.text)
        if len(b.segments) == start:
            b.warning("slide_without_text", "No extractable text; may contain meaningful visuals", loc)
    b.warning("pptx_visuals_not_fully_extracted", "Master/layout text, comments, SmartArt, charts, equations, media and spatial reading order are not fully reproduced")
    b.details["slide_count"] = len(presentation.slides)


def read_pdf(raw: bytes, b: Builder):
    if not raw.startswith(b"%PDF-"):
        raise ExtractionError("format_mismatch", "Not a PDF file")
    pypdf = b.dependency("pypdf", "pypdf")
    reader = pypdf.PdfReader(BytesIO(raw), strict=True)
    if reader.is_encrypted:
        raise ExtractionError("encrypted_document", "Decrypt an authorized copy outside this tool")
    if len(reader.pages) > b.limits.max_pages_or_slides:
        raise ExtractionError("extraction_budget_exceeded", "Too many PDF pages")
    for n, page in enumerate(reader.pages, 1):
        loc = f"PDF/page={n}"
        contents = page.get_contents()
        if contents is not None and len(contents.get_data()) > b.limits.max_pdf_stream_bytes:
            raise ExtractionError("extraction_budget_exceeded", f"PDF content stream too large at {loc}")
        text = page.extract_text() or ""
        if not text.strip():
            b.warning("pdf_page_without_text", "Page may be blank, scanned or image-only; no OCR was attempted", loc)
        b.add(loc, text)
        resources = page.get("/Resources")
        if resources is not None:
            resources = resources.get_object()
            xobjects = resources.get("/XObject")
            if xobjects is not None:
                xobjects = xobjects.get_object()
                if len(xobjects) > b.limits.max_segments:
                    raise ExtractionError("extraction_budget_exceeded", "Too many PDF XObjects")
                for name, reference in sorted(xobjects.items(), key=lambda item: str(item[0])):
                    obj = reference.get_object()
                    asset_loc = f"{loc}/xobject={name}"
                    if obj.get("/Subtype") == "/Image":
                        b.add(asset_loc, json.dumps({"name": str(name), "width": obj.get("/Width"),
                                                      "height": obj.get("/Height"), "kind": "image"},
                                                     sort_keys=True), "embedded_asset")
                        b.warning("pdf_embedded_image_not_extracted", "Image pixels and visual meaning were not extracted", asset_loc)
                    elif obj.get("/Subtype") == "/Form":
                        b.warning("pdf_form_xobject_not_inspected", "Form XObject may contain visual content; nested resources were not inspected", asset_loc)
    b.details["physical_page_count"] = len(reader.pages)
    b.warning("pdf_layout_and_visuals_require_review", "Text order/tables/figures/annotations/forms and existing OCR errors require original-page review; page numbers are physical 1-based positions")


def read_csv(raw: bytes, b: Builder, options: Options):
    if len(options.csv_delimiter) != 1 or options.csv_delimiter in "\r\n\x00\"":
        raise ExtractionError("invalid_csv_options", "Specify one delimiter character other than a quote/newline/NUL")
    try:
        text = raw.decode(options.csv_encoding, errors="strict")
    except (LookupError, UnicodeError) as exc:
        raise ExtractionError("csv_encoding_error", "CSV decoding failed; explicitly specify its encoding, do not guess") from exc
    if "\x00" in text:
        raise ExtractionError("invalid_text", "CSV contains NUL")
    if options.csv_header not in {"none", "first-row"} or (options.csv_header != "none" and not options.structured_csv):
        raise ExtractionError("invalid_csv_options", "CSV header choice requires structured CSV representation")
    b.details.update(csv_encoding=options.csv_encoding, csv_delimiter=options.csv_delimiter,
                     csv_representation="structured" if options.structured_csv else "raw-records",
                     csv_header=options.csv_header, header_inferred=False)
    cells, first_width, widths_differ = 0, None, False
    headers = None
    try:
        reader = csv.reader(StringIO(text, newline=""), delimiter=options.csv_delimiter, strict=True)
        previous_line = 0
        for n, row in enumerate(reader, 1):
            start, end = previous_line + 1, reader.line_num
            previous_line = end
            cells += len(row)
            if cells > b.limits.max_cells:
                raise ExtractionError("extraction_budget_exceeded", "Too many CSV cells")
            if first_width is None:
                first_width = len(row)
            elif len(row) != first_width:
                widths_differ = True
            if options.structured_csv and options.csv_header == "first-row" and n == 1 and not row:
                raise ExtractionError("csv_declared_header_missing", "First CSV record is empty; declared header cannot be inferred from a later row")
            if row:
                locator = f"CSV/record={n}/physical-lines={start}-{end}"
                # JSON strings preserve empty values, embedded newlines, leading zeros and literal formulas.
                if not options.structured_csv:
                    b.add(locator, json.dumps(row, ensure_ascii=False))
                elif options.csv_header == "first-row" and headers is None:
                    headers = list(row)
                    if any(not label for label in headers):
                        b.warning("csv_empty_declared_header", "At least one declared column label is empty; use column positions", locator)
                    if len(set(headers)) != len(headers):
                        b.warning("csv_duplicate_declared_header", "Declared header labels repeat; use column positions", locator)
                    b.add(locator, json.dumps({"declared_header": headers}, ensure_ascii=False), "table_header")
                else:
                    columns = [{"column": i, "header": headers[i - 1] if headers is not None and i <= len(headers) else None,
                                "value": value} for i, value in enumerate(row, 1)]
                    b.add(locator, json.dumps({"columns": columns}, ensure_ascii=False), "table_row")
    except csv.Error as exc:
        raise ExtractionError("csv_parse_error", str(exc)) from exc
    if widths_differ:
        b.warning("csv_ragged_rows", "Records have different field counts; review delimiter and source structure")
    if first_width == 1:
        b.warning("csv_single_column", "Parsed as one column; verify the explicit delimiter")
    b.warning("csv_values_not_interpreted", "Fields are strings; header, locale, dates, numbers and formulas are not inferred or evaluated")


def read_xlsx(raw: bytes, b: Builder):
    """Read stored OOXML cells directly; preserve lexical numbers and formula caches.

    This intentionally does not calculate formulas or recreate Excel display formatting.
    No workbook writer or office converter is invoked.
    """
    parts = ooxml_parts(raw, b)
    n = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    main = parts.get("xl/workbook.xml")
    rels = parts.get("xl/_rels/workbook.xml.rels")
    if main is None or rels is None:
        raise ExtractionError("format_mismatch", "Not a supported transitional XLSX package")
    targets = {}
    for rel in rels:
        if rel.get("TargetMode") == "External":
            continue
        value = rel.get("Target", "")
        if ".." in PurePosixPath(value).parts or "\\" in value:
            raise ExtractionError("invalid_container", "Unsafe workbook relationship")
        targets[rel.get("Id")] = value.lstrip("/") if value.startswith("/") else "xl/" + value
    shared = []
    string_table = parts.get("xl/sharedStrings.xml")
    if string_table is not None:
        for item in string_table.findall("s:si", n):
            # Exclude phonetic annotations (rPh) rather than appending them to a value.
            shared.append("".join(e.text or "" for e in item.findall("s:t", n) + item.findall("s:r/s:t", n)))
    date_prop = main.find("s:workbookPr", n)
    date_system = "1904" if date_prop is not None and date_prop.get("date1904") in {"1", "true"} else "1900"
    b.details["excel_date_system"] = date_system
    b.details["xlsx_value_representation"] = "Stored lexical values and formula text; formatting not rendered"
    styles_root = parts.get("xl/styles.xml")
    custom_formats, cell_styles = {}, []
    if styles_root is not None:
        custom_formats = {e.get("numFmtId"): e.get("formatCode") for e in styles_root.findall("s:numFmts/s:numFmt", n)}
        cell_styles = [e.get("numFmtId", "0") for e in styles_root.findall("s:cellXfs/s:xf", n)]
    count = 0
    sheets = main.findall("s:sheets/s:sheet", n)
    if len(sheets) > b.limits.max_pages_or_slides:
        raise ExtractionError("extraction_budget_exceeded", "Too many sheets")
    names = [sheet.get("name") for sheet in sheets]
    if len(set(names)) != len(names):
        raise ExtractionError("invalid_container", "Duplicate worksheet names")
    for sheet in sheets:
        name = sheet.get("name", "")
        loc = "XLSX/sheet=" + json.dumps(name, ensure_ascii=False)
        target = targets.get(sheet.get(rns + "id"))
        root = parts.get(target)
        if root is None:
            raise ExtractionError("invalid_container", f"Missing worksheet part for {name!r}")
        if root.tag != "{" + n["s"] + "}worksheet":
            b.warning("xlsx_nonworksheet_part", "Chart sheet or other non-cell sheet was not extracted", loc)
            continue
        state = sheet.get("state", "visible")
        if state != "visible":
            b.warning("hidden_sheet_included", f"Sheet state is {state}; stored cells are included", loc)
        hidden_columns = []
        for col in root.findall("s:cols/s:col", n):
            if col.get("hidden") in {"1", "true"}:
                hidden_columns.append((int(col.get("min", "0")), int(col.get("max", "0"))))
        seen_addresses = set()
        for row in root.findall("s:sheetData/s:row", n):
            for cell in row.findall("s:c", n):
                count += 1
                if count > b.limits.max_cells:
                    raise ExtractionError("extraction_budget_exceeded", "Too many stored XLSX cells")
                address = cell.get("r", "")
                if address in seen_addresses:
                    raise ExtractionError("invalid_container", "Duplicate worksheet cell coordinate")
                seen_addresses.add(address)
                match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", address)
                if match is None:
                    raise ExtractionError("invalid_container", "Invalid/missing XLSX cell coordinate")
                column = 0
                for c in match[1]:
                    column = column * 26 + ord(c) - 64
                if column > 16384 or int(match[2]) > 1048576:
                    raise ExtractionError("invalid_container", "XLSX cell outside worksheet limits")
                value = cell.find("s:v", n)
                formula = cell.find("s:f", n)
                inline = cell.find("s:is", n)
                kind = cell.get("t", "n")
                result = {"stored_type": kind, "sheet_state": state,
                          "row_hidden": row.get("hidden") in {"1", "true"},
                          "column_hidden": any(lo <= column <= hi for lo, hi in hidden_columns)}
                if value is None and formula is None and inline is None:
                    continue  # A style-only cell contains no source value.
                if kind == "s" and value is not None:
                    idx = int(value.text or "-1")
                    if idx < 0 or idx >= len(shared):
                        raise ExtractionError("invalid_container", "Invalid shared string index")
                    result["value"] = shared[idx]
                elif kind == "inlineStr" and inline is not None:
                    result["value"] = "".join(e.text or "" for e in inline.findall("s:t", n) + inline.findall("s:r/s:t", n))
                else:
                    result["value"] = value.text if value is not None else None
                if formula is not None:
                    result["cached_value"] = result.pop("value", None)
                    result.update(formula=formula.text, formula_attributes=dict(formula.attrib),
                                  cached_value_present=value is not None and value.text is not None,
                                  cache_status="unverified_not_recalculated")
                    if not result["cached_value_present"]:
                        b.warning("xlsx_formula_cache_missing", "At least one formula lacks a stored cached result; null is not zero")
                    if formula.attrib:
                        b.warning("xlsx_special_formula", "Shared/array formula metadata is retained, not expanded or recalculated")
                style_index = int(cell.get("s", "0"))
                if style_index < 0 or (cell_styles and style_index >= len(cell_styles)):
                    raise ExtractionError("invalid_container", "Invalid cell style index")
                numfmt = cell_styles[style_index] if cell_styles else "0"
                result.update(style_index=style_index, number_format_id=numfmt)
                if numfmt in custom_formats:
                    result["custom_number_format"] = custom_formats[numfmt]
                b.add(loc + f"/cell={address}", json.dumps(result, ensure_ascii=False, sort_keys=True))
        if root.find("s:mergeCells", n) is not None:
            b.warning("xlsx_merged_cells", "Merged ranges: " + ", ".join(e.get("ref", "") for e in root.findall("s:mergeCells/s:mergeCell", n)), loc)
    b.warning("xlsx_formula_caches_unverified", "Formula caches may be absent or stale; no formula was executed")
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        for name in sorted(item.filename for item in archive.infolist()
                           if item.filename.startswith("xl/media/") and not item.is_dir()):
            asset = archive.read(name)
            b.add("XLSX/part=" + name, json.dumps({"part": name, "bytes": len(asset),
                                                   "sha256": digest(asset)}, sort_keys=True), "embedded_asset")
            b.warning("xlsx_embedded_image_not_extracted",
                      "Workbook media part is present; visual meaning and sheet placement were not reconstructed", name)
        for name in sorted(item.filename for item in archive.infolist()
                           if item.filename.startswith("xl/charts/") and not item.is_dir()):
            b.warning("xlsx_chart_not_extracted", "Workbook chart meaning was not extracted", name)
    b.warning("xlsx_visuals_and_metadata_not_extracted", "Charts, comments, pivots, images, validations and formatted display values require original review; numeric/date serials remain stored strings")


_CHARSETS = {"utf-8": "utf-8", "utf8": "utf-8", "us-ascii": "ascii",
             "ascii": "ascii", "windows-1252": "cp1252", "cp1252": "cp1252",
             "iso-8859-1": "iso-8859-1", "latin-1": "iso-8859-1"}
_HTML_BLOCKS = {"title": "heading", "h1": "heading", "h2": "heading", "h3": "heading",
                "h4": "heading", "h5": "heading", "h6": "heading", "p": "paragraph",
                "li": "list_item", "th": "table_cell", "td": "table_cell",
                "pre": "code", "code": "code", "blockquote": "quote", "body": "paragraph",
                "div": "paragraph", "section": "paragraph", "article": "paragraph",
                "main": "paragraph", "header": "paragraph", "footer": "paragraph",
                "dt": "paragraph", "dd": "paragraph", "caption": "caption", "text": "diagram_label"}
_HTML_VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link",
              "meta", "param", "source", "track", "wbr"}
_HTML_SUPPRESSED = {"script", "style", "canvas", "iframe", "object", "embed", "template", "foreignobject"}


def _decode_declared(raw: bytes, declared: str | None, context: str) -> tuple[str, str]:
    charset = (declared or "utf-8").strip().lower()
    encoding = _CHARSETS.get(charset)
    if encoding is None:
        raise ExtractionError("unsupported_charset", f"{context}: unsupported declared charset {charset!r}")
    if raw.startswith(b"\xef\xbb\xbf"):
        if encoding != "utf-8":
            raise ExtractionError("charset_conflict", f"{context}: UTF-8 BOM conflicts with declared charset")
        raw = raw[3:]
    try:
        return raw.decode(encoding, errors="strict"), charset
    except UnicodeError as exc:
        raise ExtractionError("charset_decode_error", f"{context}: bytes disagree with declared charset") from exc


class _InertHTML(HTMLParser):
    """Bounded structure capture; source markup never runs or fetches anything."""

    def __init__(self, b: Builder, prefix: str):
        super().__init__(convert_charrefs=True)
        self.b, self.prefix = b, prefix
        self.stack = []
        self.roots = {}
        self.nodes = 0
        self.cids = set()
        self.malformed = False
        self.root_text = []
        self.last_heading = None
        self.cells = 0

    def handle_starttag(self, tag, attrs):
        self.nodes += 1
        if self.nodes > self.b.limits.max_structure_nodes or len(self.stack) >= self.b.limits.max_structure_depth:
            raise ExtractionError("extraction_budget_exceeded", "HTML structure exceeds node/depth budget")
        parent = self.stack[-1] if self.stack else None
        siblings = parent["children"] if parent else self.roots
        siblings[tag] = siblings.get(tag, 0) + 1
        loc = (parent["path"] if parent else self.prefix) + f"/{tag}={siblings[tag]}"
        attributes = dict(attrs)
        if tag in _HTML_SUPPRESSED:
            self.b.warning("html_active_content_not_extracted" if tag in {"script", "style", "iframe", "object", "embed"}
                           else "html_visual_content_not_extracted",
                           f"{tag} content is not interpreted", loc)
        if tag == "svg" and not (parent and parent["suppressed"]):
            self.b.warning("html_visual_content_not_extracted", "Inline SVG labels and geometry were captured, but rendering and meaning were not interpreted", loc)
            self.b.add(loc + "/canvas", json.dumps({key: attributes.get(key) for key in ("viewbox", "width", "height")},
                                                   ensure_ascii=False, sort_keys=True), "image_region")
        if not (parent and parent["suppressed"]) and any(frame["tag"] == "svg" for frame in self.stack) and tag in _SVG_GEOMETRY:
            self.b.add(loc + "/geometry", json.dumps({"tag": tag, "attributes": attributes},
                                                     ensure_ascii=False, sort_keys=True), "vector_geometry")
            if tag == "image":
                self.b.warning("html_image_not_extracted", "Inline SVG referenced image was not decoded or interpreted", loc)
        if any(key.startswith("on") and value for key, value in attrs) or \
                any((attributes.get(key) or "").strip().lower().startswith("javascript:")
                    for key in ("src", "href", "data")):
            self.b.warning("html_active_content_not_extracted", "Event handler or JavaScript URL was not executed", loc)
        # The enclosing unavailable element keeps its own references. Its
        # suppressed descendants cannot supply independent citable content.
        for key in (() if parent and parent["suppressed"] else ("src", "href", "data", "poster")):
            value = attributes.get(key) or ""
            if value:
                self.b.add(loc + f"/attribute={key}",
                           json.dumps({"tag": tag, "attribute": key, "value": value}, ensure_ascii=False),
                           "resource_reference")
            if value.startswith("cid:"):
                self.cids.add(value[4:].strip("<>"))
            elif value and (value.startswith(("http:", "https:", "//", "data:")) or tag in {"img", "link", "iframe", "object", "embed", "source"}):
                self.b.warning("html_external_resource_not_fetched", "Referenced resource was not fetched or interpreted", loc)
        if tag == "img" and not (parent and parent["suppressed"]):
            self.b.warning("html_image_not_extracted", "Image pixels and visual meaning were not extracted", loc)
            if attributes.get("alt"):
                self.b.add(loc + "/alt", attributes["alt"], "image_alt")
        if tag == "br":
            self.handle_data("\n")
        if tag not in _HTML_VOID:
            self.stack.append({"tag": tag, "path": loc, "children": {}, "text": [],
                               "suppressed": (parent and parent["suppressed"]) or tag in _HTML_SUPPRESSED,
                               "attrs": attributes, "cells": [], "headers": [], "caption": None,
                               "rows_seen": 0, "spans_require_review": False})

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in _HTML_VOID:
            self.handle_endtag(tag)

    def handle_data(self, data):
        if self.stack and self.stack[-1]["suppressed"]:
            return
        cell = next((frame for frame in reversed(self.stack) if frame["tag"] in {"th", "td"}), None)
        if cell is not None:
            cell["text"].append(data)
        for frame in reversed(self.stack):
            if frame["tag"] in _HTML_BLOCKS:
                if frame is not cell:
                    frame["text"].append(data)
                # Inline code is also part of the containing assertion. Keep
                # its own locator without dropping qualifiers from the parent.
                if frame["tag"] != "code":
                    return
        self.root_text.append(data)

    def handle_endtag(self, tag):
        if not any(frame["tag"] == tag for frame in self.stack):
            self.malformed = True
            return
        while self.stack:
            frame = self.stack.pop()
            if frame["tag"] in _HTML_BLOCKS and not frame["suppressed"]:
                value = "".join(frame["text"])
                self.b.add(frame["path"], value, _HTML_BLOCKS[frame["tag"]])
                if frame["tag"] in {"h1", "h2", "h3", "h4", "h5", "h6"} and value.strip():
                    self.last_heading = value.strip()
                if frame["tag"] == "caption":
                    table = next((item for item in reversed(self.stack) if item["tag"] == "table"), None)
                    if table is not None:
                        table["caption"] = value.strip() or None
            if frame["tag"] in {"th", "td"} and not frame["suppressed"]:
                self.cells += 1
                if self.cells > self.b.limits.max_cells:
                    raise ExtractionError("extraction_budget_exceeded", "HTML table cell count exceeds budget")
                row = next((item for item in reversed(self.stack) if item["tag"] == "tr"), None)
                if row is not None:
                    row["cells"].append({"tag": frame["tag"], "value": "".join(frame["text"]).strip(),
                                         "colspan": frame["attrs"].get("colspan", "1"),
                                         "rowspan": frame["attrs"].get("rowspan", "1")})
            if frame["tag"] == "tr" and frame["cells"] and not frame["suppressed"]:
                table = next((item for item in reversed(self.stack) if item["tag"] == "table"), None)
                cells = frame["cells"]
                if table is None:
                    self.b.warning("html_malformed_structure", "Table row has no enclosing table", frame["path"])
                else:
                    table["rows_seen"] += 1
                has_spans = any(cell["colspan"] != "1" or cell["rowspan"] != "1" for cell in cells)
                if table is not None and has_spans:
                    # Without reconstructing the grid, merged headers and
                    # inherited rowspans make later alignment unreliable too.
                    table["spans_require_review"] = True
                unreliable_alignment = has_spans or (table is not None and table["spans_require_review"])
                if unreliable_alignment:
                    self.b.warning("html_table_spans_require_review", "Merged cells prevent reliable header alignment", frame["path"])
                if all(cell["tag"] == "th" for cell in cells) and table is not None and table["rows_seen"] == 1:
                    table["headers"] = [cell["value"] for cell in cells]
                else:
                    if all(cell["tag"] == "th" for cell in cells) and table is not None:
                        self.b.warning("html_additional_header_row", "Additional header-only row is preserved as a row; multi-level header meaning requires review", frame["path"])
                    headers = table["headers"] if table is not None else []
                    if headers and len(headers) != len(cells) and not unreliable_alignment:
                        self.b.warning("html_table_width_mismatch", "Data row width differs from the declared header row", frame["path"])
                    columns = [{"column": n, "header": headers[n - 1] if not unreliable_alignment and n <= len(headers) else None,
                                **cell} for n, cell in enumerate(cells, 1)]
                    self.b.add(frame["path"] + "/row", json.dumps({"preceding_heading": self.last_heading,
                               "table_caption": table["caption"] if table is not None else None,
                               "columns": columns}, ensure_ascii=False, sort_keys=True), "table_row")
            if frame["tag"] == tag:
                break
            self.malformed = True

    def finish(self):
        if self.stack:
            self.malformed = True
            while self.stack:
                self.handle_endtag(self.stack[-1]["tag"])
        if self.malformed:
            self.b.warning("html_malformed_structure", "Unmatched or unclosed tags; review source structure", self.prefix)
        self.b.add(self.prefix + "/text", "".join(self.root_text), "paragraph")


def _read_html_text(text: str, b: Builder, prefix: str) -> set[str]:
    parser = _InertHTML(b, prefix)
    try:
        parser.feed(text)
        parser.close()
    except ExtractionError:
        raise
    except (ValueError, AssertionError) as exc:
        raise ExtractionError("html_parse_error", "HTML tokenizer rejected malformed source") from exc
    parser.finish()
    return parser.cids


def read_html(raw: bytes, b: Builder):
    head = raw[:4096].decode("ascii", errors="ignore")
    match = re.search(r"<meta\s+[^>]*charset\s*=\s*['\"]?([\w-]+)", head, re.I)
    declared = match.group(1) if match else None
    text, charset = _decode_declared(raw, declared, "HTML")
    b.details.update(declared_charset=declared, decoded_charset=charset)
    _read_html_text(text, b, "HTML")
    b.warning("html_dynamic_content_unavailable", "Scripts, rendered state and embedded visual semantics are not captured")


def _email_payload_bytes(part, locator: str) -> bytes:
    transfer = (part.get("Content-Transfer-Encoding") or "").lower()
    if transfer == "base64":
        payload = part.get_payload()
        try:
            return base64.b64decode(re.sub(r"\s+", "", payload), validate=True)
        except (ValueError, base64.binascii.Error) as exc:
            raise ExtractionError("eml_transfer_decode_error", f"Invalid base64 at {locator}") from exc
    result = part.get_payload(decode=True)
    if result is None:
        payload = part.get_payload()
        if not isinstance(payload, str):
            raise ExtractionError("eml_payload_error", f"Unexpected payload at {locator}")
        return payload.encode("ascii", errors="surrogateescape")
    return result


def read_eml(raw: bytes, b: Builder):
    message = BytesParser(policy=policy.default).parsebytes(raw)
    if message.defects:
        raise ExtractionError("eml_malformed_structure", "Email headers or MIME boundaries are malformed")
    b.parsers["stdlib-email"] = "policy.default"
    b.details["date_header_unverified"] = str(message.get("Date", ""))
    for field in ("From", "To", "Cc", "Date", "Subject", "Message-ID", "In-Reply-To", "References"):
        for n, value in enumerate(message.get_all(field, []), 1):
            b.add(f"EML/header/{field.lower()}={n}", f"{field}: {value}", "email_header")
    parts = 0
    content_ids = set()
    referenced_ids = set()

    def walk(part, loc, depth):
        nonlocal parts
        parts += 1
        if parts > b.limits.max_mime_parts or depth > b.limits.max_structure_depth:
            raise ExtractionError("extraction_budget_exceeded", "MIME part count/depth exceeds budget")
        if part.defects:
            raise ExtractionError("eml_malformed_structure", f"Malformed MIME part at {loc}")
        media = part.get_content_type()
        disposition = part.get_content_disposition()
        filename = part.get_filename()
        cid = (part.get("Content-ID") or "").strip("<>")
        if cid:
            content_ids.add(cid)
        if disposition == "attachment" or filename or media == "message/rfc822":
            compound = part.is_multipart()
            payload = part.as_bytes(policy=policy.default) if compound else _email_payload_bytes(part, loc)
            if len(payload) > b.limits.max_member_bytes:
                raise ExtractionError("extraction_budget_exceeded", "MIME attachment exceeds byte budget")
            inventory = {"filename": filename, "media_type": media, "disposition": disposition,
                         "content_id": cid or None, "bytes": len(payload), "sha256": digest(payload),
                         "hash_basis": "serialized_mime_part" if compound else "decoded_payload"}
            b.add(loc + "/attachment", json.dumps(inventory, ensure_ascii=False, sort_keys=True),
                  "attachment_metadata")
            b.warning("eml_attachment_not_extracted", "Attachment content is separately scoped and was not extracted", loc)
            return
        if part.is_multipart():
            for n, child in enumerate(part.iter_parts(), 1):
                walk(child, f"{loc}/part={n}", depth + 1)
            return
        payload = _email_payload_bytes(part, loc)
        if len(payload) > b.limits.max_member_bytes:
            raise ExtractionError("extraction_budget_exceeded", "MIME part exceeds byte budget")
        if media not in {"text/plain", "text/html"}:
            inventory = {"filename": None, "media_type": media, "disposition": disposition,
                         "content_id": cid or None, "bytes": len(payload), "sha256": digest(payload),
                         "hash_basis": "decoded_payload"}
            b.add(loc + "/attachment", json.dumps(inventory, ensure_ascii=False, sort_keys=True),
                  "attachment_metadata")
            b.warning("eml_attachment_not_extracted", "Non-text part content was not extracted", loc)
            return
        text, charset = _decode_declared(payload, part.get_content_charset(), loc)
        b.details.setdefault("body_charsets", []).append({"locator": loc, "declared": part.get_content_charset(),
                                                            "decoded": charset})
        if media == "text/html":
            referenced_ids.update(_read_html_text(text, b, loc + "/html"))
        else:
            normal, quote, block = [], [], 0
            for line in text.splitlines():
                if line.lstrip().startswith(">"):
                    if normal:
                        block += 1
                        b.add(f"{loc}/plain={block}", "\n".join(normal), "email_body")
                        normal = []
                    quote.append(line)
                else:
                    if quote:
                        block += 1
                        b.add(f"{loc}/quoted-reply={block}", "\n".join(quote), "email_quote")
                        quote = []
                    normal.append(line)
            if normal:
                block += 1
                b.add(f"{loc}/plain={block}", "\n".join(normal), "email_body")
            if quote:
                block += 1
                b.add(f"{loc}/quoted-reply={block}", "\n".join(quote), "email_quote")

    walk(message, "EML", 0)
    for cid in sorted(referenced_ids - content_ids):
        b.warning("eml_referenced_attachment_missing", "Referenced content-id is absent from this message", f"cid:{cid}")
    b.details["mime_part_count"] = parts


_SVG_NS = "http://www.w3.org/2000/svg"
_SVG_GEOMETRY = {"rect", "circle", "ellipse", "line", "polyline", "polygon", "path", "image", "use"}


def _svg_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _svg_attributes(element) -> dict:
    result = {}
    for key, value in element.attrib.items():
        if key.startswith("{http://www.w3.org/1999/xlink}"):
            key = "xlink:" + _svg_name(key)
        else:
            key = _svg_name(key)
        result[key] = value
    return result


def _svg_label_text(element, max_depth: int, depth: int = 0) -> str:
    if depth > max_depth:
        raise ExtractionError("extraction_budget_exceeded", "SVG label nesting exceeds budget")
    pieces = [element.text or ""]
    for child in element:
        if _svg_name(child.tag) == "tspan":
            pieces.append(_svg_label_text(child, max_depth, depth + 1))
        pieces.append(child.tail or "")
    return "".join(pieces)


def read_svg(raw: bytes, b: Builder):
    """Preserve explicit XML labels and vector primitives without rendering them."""
    safe_xml = b.dependency("defusedxml.ElementTree", "defusedxml")
    try:
        root = safe_xml.fromstring(raw, forbid_dtd=True)
    except Exception as exc:
        raise ExtractionError("unsafe_or_invalid_xml", "SVG XML is invalid or contains prohibited declarations/entities") from exc
    if root.tag != "{" + _SVG_NS + "}svg":
        raise ExtractionError("format_mismatch", "Expected an SVG root element in the SVG namespace")
    canvas = _svg_attributes(root)
    b.add("SVG/canvas", json.dumps({"viewBox": canvas.get("viewBox"), "width": canvas.get("width"),
                                     "height": canvas.get("height")}, ensure_ascii=False, sort_keys=True),
          "image_region")
    nodes = 0
    stack = [(root, "SVG/svg=1", 0, False)]
    while stack:
        element, loc, depth, in_defs = stack.pop()
        nodes += 1
        if nodes > b.limits.max_structure_nodes or depth > b.limits.max_structure_depth:
            raise ExtractionError("extraction_budget_exceeded", "SVG node/depth budget exceeded")
        tag = _svg_name(element.tag)
        attributes = _svg_attributes(element)
        in_defs = in_defs or tag == "defs"
        if tag in {"script", "style", "foreignObject"}:
            b.warning("svg_active_or_foreign_content_not_extracted", f"{tag} content was not executed or interpreted", loc)
            continue
        if tag in {"title", "desc", "text"}:
            value = _svg_label_text(element, b.limits.max_structure_depth).strip()
            if value:
                b.add(loc, value, "diagram_label" if tag == "text" else "diagram_caption")
            if tag == "text" and attributes:
                b.add(loc + "/attributes", json.dumps({"tag": tag, "attributes": attributes,
                                                         "inside_definitions": in_defs},
                                                        ensure_ascii=False, sort_keys=True), "vector_geometry")
        if tag in _SVG_GEOMETRY:
            b.add(loc + "/geometry", json.dumps({"tag": tag, "attributes": attributes,
                                                  "inside_definitions": in_defs},
                                                 ensure_ascii=False, sort_keys=True), "vector_geometry")
        if tag == "image":
            b.warning("svg_embedded_image_not_extracted", "Referenced image pixels and meaning were not extracted", loc)
        if tag == "use":
            b.warning("svg_use_not_expanded", "Referenced symbol/use instance was not resolved or rendered", loc)
        href = attributes.get("href") or attributes.get("xlink:href")
        if href and not href.startswith("#"):
            b.warning("svg_external_resource_not_fetched", "Referenced resource was not fetched or decoded", loc)
        if tag in {"filter", "mask", "clipPath", "pattern", "animate", "animateTransform"}:
            b.warning("svg_render_effect_not_interpreted", f"{tag} rendering effect was not interpreted", loc)
        children = list(element)
        for index, child in reversed(list(enumerate(children, 1))):
            stack.append((child, loc + f"/{_svg_name(child.tag)}={index}", depth + 1, in_defs))
    b.details["svg_node_count"] = nodes
    b.details["svg_viewbox_source"] = canvas.get("viewBox")
    b.warning("svg_rendering_unverified", "Stored XML labels and geometry may be hidden, transformed, illegible or contradicted by the rendered diagram")
    b.warning("svg_visual_meaning_unavailable", "No rendered-image interpretation or node/edge/cardinality inference was performed")


def _png_scanline_lengths(width: int, height: int, bit_depth: int,
                          color_type: int, interlace: int):
    """Yield encoded row lengths, excluding each row's filter byte."""
    channels = {0: 1, 2: 3, 3: 1, 4: 2, 6: 4}[color_type]
    passes = [(0, 0, 1, 1)] if interlace == 0 else [
        (0, 0, 8, 8), (4, 0, 8, 8), (0, 4, 4, 8), (2, 0, 4, 4),
        (0, 2, 2, 4), (1, 0, 2, 2), (0, 1, 1, 2)]
    for x, y, dx, dy in passes:
        if width <= x or height <= y:
            continue
        columns = (width - x + dx - 1) // dx
        rows = (height - y + dy - 1) // dy
        size = (columns * channels * bit_depth + 7) // 8
        for _ in range(rows):
            yield size


def _validate_png_idat(parts: list[bytes], width: int, height: int,
                       bit_depth: int, color_type: int, interlace: int):
    """Bounded zlib/scanline validation, without reconstructing pixel values."""
    lengths = _png_scanline_lengths(width, height, bit_depth, color_type, interlace)
    # A PNG row has one filter byte followed by its packed sample bytes.
    remaining = 0
    completed = 0
    inflater = zlib.decompressobj()
    try:
        for part in parts:
            if inflater.eof and part:
                raise ExtractionError("invalid_png", "PNG IDAT contains trailing compressed bytes")
            pending = part
            while True:
                decoded = inflater.decompress(pending, 64 * 1024)
                pending = inflater.unconsumed_tail
                offset = 0
                while offset < len(decoded):
                    if remaining == 0:
                        try:
                            remaining = next(lengths)
                        except StopIteration as exc:
                            raise ExtractionError("extraction_budget_exceeded", "PNG pixel stream exceeds declared dimensions") from exc
                        if decoded[offset] > 4:
                            raise ExtractionError("invalid_png", "PNG scanline has an invalid filter type")
                        completed += 1
                        offset += 1
                    else:
                        take = min(remaining, len(decoded) - offset)
                        remaining -= take
                        offset += take
                if inflater.eof:
                    if inflater.unused_data or pending:
                        raise ExtractionError("invalid_png", "PNG IDAT has trailing compressed bytes")
                    break
                if not pending and len(decoded) < 64 * 1024:
                    break
        if not inflater.eof or remaining or next(lengths, None) is not None:
            raise ExtractionError("invalid_png", "PNG pixel stream is truncated or disagrees with dimensions")
    except zlib.error as exc:
        raise ExtractionError("invalid_png", "PNG IDAT zlib stream is invalid") from exc
    return completed


def read_png(raw: bytes, b: Builder):
    """Validate PNG envelope and packed scanline structure; never infer pixels."""
    if not raw.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ExtractionError("format_mismatch", "Expected PNG signature")
    offset, chunks, seen_idat, seen_iend = 8, 0, False, False
    idat_parts, idat_closed, seen_plte = [], False, False
    width = height = bit_depth = color_type = interlace = None
    while offset < len(raw):
        if offset + 12 > len(raw):
            raise ExtractionError("invalid_png", "Truncated PNG chunk")
        length = struct.unpack_from(">I", raw, offset)[0]
        kind = raw[offset + 4:offset + 8]
        end = offset + 12 + length
        if end > len(raw):
            raise ExtractionError("invalid_png", "PNG chunk exceeds source bytes")
        data = raw[offset + 8:offset + 8 + length]
        expected_crc = struct.unpack_from(">I", raw, offset + 8 + length)[0]
        if zlib.crc32(kind + data) & 0xffffffff != expected_crc:
            raise ExtractionError("invalid_png", "PNG chunk CRC mismatch")
        chunks += 1
        if chunks > b.limits.max_png_chunks:
            raise ExtractionError("extraction_budget_exceeded", "PNG chunk count exceeds budget")
        if chunks == 1 and kind != b"IHDR":
            raise ExtractionError("invalid_png", "PNG IHDR must be first")
        if seen_idat and kind != b"IDAT":
            idat_closed = True
        if kind == b"IHDR":
            if chunks != 1 or length != 13:
                raise ExtractionError("invalid_png", "Duplicate or malformed PNG IHDR")
            width, height, bit_depth, color_type, compression, filtering, interlace = struct.unpack(">IIBBBBB", data)
            valid_depths = {0: {1, 2, 4, 8, 16}, 2: {8, 16}, 3: {1, 2, 4, 8},
                            4: {8, 16}, 6: {8, 16}}
            if width == 0 or height == 0 or width * height > b.limits.max_image_pixels:
                raise ExtractionError("extraction_budget_exceeded", "PNG pixel dimensions exceed budget")
            if bit_depth not in valid_depths.get(color_type, set()) or compression != 0 or filtering != 0 or interlace not in {0, 1}:
                raise ExtractionError("invalid_png", "Unsupported or invalid PNG IHDR values")
        elif kind == b"IDAT":
            if idat_closed:
                raise ExtractionError("invalid_png", "PNG IDAT chunks must be consecutive")
            seen_idat = True
            idat_parts.append(data)
        elif kind == b"PLTE":
            if seen_plte or seen_idat or color_type in {0, 4} or not length or length % 3 or length > 768 \
                    or (color_type == 3 and length // 3 > 2 ** bit_depth):
                raise ExtractionError("invalid_png", "PNG palette is malformed or out of order")
            seen_plte = True
        elif kind == b"tEXt":
            key, sep, value = data.partition(b"\x00")
            if not sep or not 1 <= len(key) <= 79:
                raise ExtractionError("invalid_png", "Malformed PNG text chunk")
            b.add(f"PNG/chunk=tEXt={chunks}", json.dumps({"keyword": key.decode("latin-1"),
                  "annotation": value.decode("latin-1")}, ensure_ascii=False, sort_keys=True), "image_annotation")
            b.warning("png_annotation_unverified", "PNG text annotation is source metadata, not verified pixel meaning", f"PNG/chunk=tEXt={chunks}")
        elif kind == b"IEND":
            if length != 0 or not seen_idat or end != len(raw) or (color_type == 3 and not seen_plte):
                raise ExtractionError("invalid_png", "PNG IEND is malformed, premature or followed by bytes")
            seen_iend = True
            break
        elif kind[0] & 32 == 0:
            raise ExtractionError("invalid_png", "Unknown critical PNG chunk")
        offset = end
    if not seen_iend:
        raise ExtractionError("invalid_png", "PNG is missing IEND")
    scanlines = _validate_png_idat(idat_parts, width, height, bit_depth, color_type, interlace)
    b.add("PNG/canvas", json.dumps({"width": width, "height": height,
                                     "region_xyxy": [0, 0, width, height],
                                     "bit_depth": bit_depth, "color_type": color_type,
                                     "interlace": interlace}, sort_keys=True), "image_region")
    b.details.update(png_width=width, png_height=height, png_chunk_count=chunks,
                     png_scanlines_validated=scanlines, pixel_values_reconstructed=False)
    b.warning("png_pixels_not_interpreted", "Compressed rows were validated, but pixel values were not reconstructed, read as text, or interpreted; image meaning is unavailable", "PNG/canvas")


def read_legacy_ppt(raw: bytes, b: Builder, options: Options):
    if not raw.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        raise ExtractionError("format_mismatch", "Expected a legacy OLE .ppt, not a renamed PPTX")
    if not options.allow_legacy_ppt:
        raise ExtractionError("legacy_ppt_opt_in_required", "Use --allow-legacy-ppt only for trusted files with approved local LibreOffice, or export a reviewed PPTX")
    executable = options.soffice or shutil.which("soffice") or shutil.which("libreoffice")
    mac_path = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if executable is None and Path(mac_path).is_file():
        executable = mac_path
    if executable is None:
        raise ExtractionError("converter_missing", "LibreOffice/soffice not found; specify --soffice")
    with tempfile.TemporaryDirectory(prefix="kb-ppt-") as directory:
        root = Path(directory)
        profile = root / "profile"
        (profile / "user").mkdir(parents=True)
        # Defense in depth, not an OS sandbox or a claim of total macro/network isolation.
        (profile / "user/registrymodifications.xcu").write_text(
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
            '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
            '</item></oor:items>', encoding="utf-8")
        source = root / "source.ppt"
        source.write_bytes(raw)
        output = root / "converted"
        output.mkdir()
        try:
            ver = subprocess.run([executable, "--version"], capture_output=True, text=True, timeout=10, check=True)
            b.parsers["LibreOffice"] = ver.stdout.strip()
            command = [executable, "-env:UserInstallation=" + profile.as_uri(), "--headless", "--nologo", "--nodefault", "--norestore",
                       "--convert-to", "pptx:Impress MS PowerPoint 2007 XML", "--outdir", str(output), str(source)]
            result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                    timeout=b.limits.conversion_timeout_seconds, check=False)
        except (OSError, subprocess.SubprocessError) as exc:
            raise ExtractionError("ppt_conversion_failed", "LibreOffice launch, version check or conversion timed out/failed") from exc
        target = output / "source.pptx"
        if result.returncode != 0 or not target.is_file() or target.is_symlink():
            raise ExtractionError("ppt_conversion_failed", "LibreOffice did not produce a PPTX")
        if target.stat().st_size > b.limits.max_input_bytes:
            raise ExtractionError("extraction_budget_exceeded", "Converted PPTX exceeds input budget")
        converted = target.read_bytes()
        b.details["intermediate_pptx_sha256"] = digest(converted)
        b.warning("legacy_ppt_conversion_requires_review", "Locators refer to converted PPTX; conversion fidelity must be checked against the original PPT")
        read_pptx(converted, b, "PPT-converted-to-PPTX")


def extract_bytes(raw: bytes, filename: str, limits: Limits | None = None, options: Options | None = None) -> Extraction:
    limits, options = limits or Limits(), options or Options()
    if not isinstance(raw, bytes):
        raise TypeError("raw must be bytes from an already captured file")
    if not filename or Path(filename).name != filename or "\\" in filename or "\n" in filename or "\r" in filename:
        raise ExtractionError("invalid_filename", "Supply a basename, not a path")
    ext = Path(filename).suffix.lower()
    if ext not in FORMATS:
        raise ExtractionError("unsupported_extension", f"Supported: {', '.join(sorted(FORMATS))}")
    if not raw:
        raise ExtractionError("empty_file", "Document is zero bytes")
    if len(raw) > limits.max_input_bytes:
        raise ExtractionError("input_budget_exceeded", "Document exceeds raw-byte budget")
    b = Builder(filename, raw, limits)
    try:
        if ext == ".csv":
            read_csv(raw, b, options)
        elif ext == ".ppt":
            read_legacy_ppt(raw, b, options)
        else:
            {".docx": read_docx, ".pptx": read_pptx, ".pdf": read_pdf, ".xlsx": read_xlsx,
             ".html": read_html, ".eml": read_eml, ".svg": read_svg, ".png": read_png}[ext](raw, b)
        return b.finish(raw, ext)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("document_parse_failed", f"{type(exc).__name__}: {str(exc)[:250]}") from exc
